"""Excel maintenance of partner fields. Cases, files and business history stay untouched."""
from io import BytesIO
from typing import Literal
import re
import uuid

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pydantic import ValidationError

from .auth import record_audit
from .business_taxonomy import INDUSTRIES, REGION_TYPES, classify, project_partner, tokens
from .database import get_db
from .enablement import fail, now
from .models import PartnerCreate
from .resource_transfer import MAX_ROWS, append_text, read_workbook, text

COLUMNS = {
    'id': '伙伴ID（新增可留空）', 'name': '伙伴名称', 'intro': '简介', 'capabilities': '能力标签',
    'industries': '行业（标准值）', 'service_areas': '区域（标准值）',
    'pending_industries': '待确认行业', 'pending_regions': '待确认区域',
    'status': '状态', 'ai_profile': '完整画像', 'profile_pending': '画像待更新',
    'created_at': '创建时间（参考）', 'updated_at': '更新时间（参考）', 'profile_updated_at': '画像更新时间（参考）',
}
LONG_FIELDS = ('intro', 'capabilities', 'industries', 'service_areas', 'pending_industries', 'pending_regions', 'ai_profile')
CONTINUATION = ['伙伴ID（新增可留空）', '伙伴名称', '字段', '段号（从2开始）', '续文']
CHUNK = 32000  # Excel silently truncates cells beyond 32767 characters.


class PartnerImport(PartnerCreate):
    status: Literal['active', 'disabled']
    ai_profile: str | None = None


def export_workbook(*, template=False):
    book = Workbook()
    sheet = book.active
    sheet.title = '伙伴信息'
    append_text(sheet, list(COLUMNS.values()))
    continuation = book.create_sheet('长文本续文')
    append_text(continuation, CONTINUATION)
    if not template:
        with get_db() as conn:
            for row in conn.execute('SELECT * FROM partners ORDER BY created_at,id'):
                data = project_partner(dict(row))
                data.update(pending_industries=','.join(data['classification_pending']['industries']),
                            pending_regions=','.join(data['classification_pending']['regions']),
                            status='启用' if row['status'] == 'active' else '停用',
                            profile_pending='是' if row['materials_revision'] != row['profile_materials_revision'] else '否')
                for field in LONG_FIELDS:
                    value = data.get(field) or ''
                    data[field] = value[:CHUNK]
                    for offset in range(CHUNK, len(value), CHUNK):
                        append_text(continuation, [data['id'], data['name'], COLUMNS[field], offset // CHUNK + 1, value[offset:offset + CHUNK]])
                append_text(sheet, [data.get(key) for key in COLUMNS])
    dictionary = book.create_sheet('行业区域参考')
    append_text(dictionary, ['分类', '标准值'])
    for kind, values in [('行业', INDUSTRIES), ('国内区域', REGION_TYPES['domestic']), ('海外区域', REGION_TYPES['overseas'])]:
        for value in values:
            append_text(dictionary, [kind, value])
    notes = book.create_sheet('使用说明')
    for line in [
        '伙伴信息每行一家；名称必填，状态填启用/停用（留空默认启用），画像待更新填是/否（留空默认否）。',
        'ID相同则更新；ID未匹配时按唯一的完整公司名称匹配，仍未匹配才新增。不合并、不删除任何已有伙伴。',
        '导入按表格更新基础信息和完整画像；空白可选字段表示清空，请先导出留存再编辑。时间列仅供参考，由目标系统记录本次更新时间。',
        '行业和区域从参考表选择，多选用逗号分隔；待确认列只保留未确认的旧分类，不作为正式分类。已有待确认内容保留。',
        '长文本超过32000字时自动分到“长文本续文”，导入按伙伴ID（无ID用名称）、字段和段号拼接，不截断。请保留续文工作表。',
        '模板中不带现有伙伴数据。导出包括启用和停用伙伴；案例、附件原件、任务与历史记录不在此表格中。',
        '画像直接采用表格文字，不调用模型、不改正式能力标签字典。导入失败整批回滚，并提示错误行号。',
        '只支持.xlsx，最多2000家、10MB；使用普通单元格值，不支持公式。',
    ]:
        append_text(notes, [line])
    for ws in book:
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='C7000B')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 26
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        if ws != notes:
            ws.auto_filter.ref = ws.dimensions
    notes.column_dimensions['A'].width = 110
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def read_rows(content):
    sheets = read_workbook(content, [('伙伴信息', list(COLUMNS.values()), MAX_ROWS), ('长文本续文', CONTINUATION, 10000)])
    rows = [(line, dict(zip(COLUMNS, values))) for line, values in sheets['伙伴信息']]
    if not rows:
        fail(422, '表格中没有伙伴信息')
    def key(data):
        return ('id', text(data['id'])) if text(data['id']) else ('name', text(data['name']))
    by_key = {}
    for line, data in rows:
        if key(data) in by_key:
            fail(422, f'“伙伴信息”第{line}行：伙伴重复')
        by_key[key(data)] = data
    parts = {}
    field_names = {COLUMNS[field]: field for field in LONG_FIELDS}
    for line, (partner_id, name, label, number, value) in sheets['长文本续文']:
        row_key = key({'id': partner_id, 'name': name})
        field = field_names.get(text(label))
        if row_key not in by_key or not field or type(number) is not int or number < 2 or not isinstance(value, str):
            fail(422, f'“长文本续文”第{line}行：请检查伙伴、字段、段号和正文')
        segments = parts.setdefault((row_key, field), {})
        if number in segments:
            fail(422, f'“长文本续文”第{line}行：段号重复')
        segments[number] = value
    for (row_key, field), segments in parts.items():
        if sorted(segments) != list(range(2, len(segments) + 2)):
            fail(422, '“长文本续文”段号不连续，请保留从2开始的全部续文')
        first = by_key[row_key][field]
        by_key[row_key][field] = ('' if first is None else str(first)) + ''.join(segments[i] for i in sorted(segments))
    return rows


def import_workbook(content, actor):
    rows = read_rows(content)
    result = {'created': 0, 'updated': 0, 'unchanged': 0}
    seen = set()
    with get_db() as conn:
        conn.lock_writer()
        for line, data in rows:
            try:
                source_id, name = text(data['id']), text(data['name'])
                if source_id and not re.fullmatch(r'[A-Za-z0-9_-]{1,100}', source_id):
                    fail(422, '伙伴ID无效')
                current = conn.execute('SELECT * FROM partners WHERE id=?', (source_id,)).fetchone() if source_id else None
                names = conn.execute('SELECT * FROM partners WHERE name=?', (name,)).fetchall()
                if len(names) > 1 or (current and any(r['id'] != current['id'] for r in names)):
                    fail(409, '公司名称对应多条伙伴或与其他伙伴冲突，请先核对名称与ID')
                if current is None and names:
                    current = names[0]
                partner_id = current['id'] if current else source_id or str(uuid.uuid4())
                if partner_id in seen:
                    fail(422, '同一家伙伴只能导入一行')
                seen.add(partner_id)
                payload = PartnerImport(name=name, status={'启用': 'active', '停用': 'disabled', '': 'active'}.get(text(data['status']), ''),
                    **{field: None if data[field] in (None, '') else str(data[field]) for field in ('intro', 'capabilities', 'industries', 'service_areas', 'ai_profile')})
                fields = {field: getattr(payload, field) for field in ('name', 'intro', 'capabilities', 'industries', 'service_areas', 'status', 'ai_profile')}
                for column, field, kind in [('pending_industries', 'industries', 'industry'), ('pending_regions', 'service_areas', 'region')]:
                    pending = tokens(data[column])
                    if any(classify(value, kind)[0] for value in pending):
                        fail(422, '待确认列不能填写已确认的标准分类，请填写到标准列')
                    old_pending = classify(current[field], kind)[1] if current else []
                    combined = list(dict.fromkeys(tokens(fields[field]) + old_pending + pending))
                    fields[field] = ','.join(combined) or None
                if text(data['profile_pending']) not in ('', '是', '否'):
                    fail(422, '画像待更新须填写是或否')
                pending = text(data['profile_pending']) == '是'
                if current and all((current[field] or '') == (value or '') for field, value in fields.items()) and pending == (current['materials_revision'] != current['profile_materials_revision']):
                    result['unchanged'] += 1
                    continue
                stamp = now()
                material_revision = current['materials_revision'] if current else 0
                profile_revision = current['profile_materials_revision'] if current else 0
                if pending and material_revision == profile_revision:
                    material_revision += 1
                if not pending:
                    profile_revision = material_revision
                fields.update(updated_at=stamp, profile_updated_at=stamp, materials_revision=material_revision, profile_materials_revision=profile_revision)
                if current:
                    conn.execute(f'UPDATE partners SET {",".join(field + "=?" for field in fields)} WHERE id=?', (*fields.values(), partner_id))
                    result['updated'] += 1
                else:
                    fields.update(id=partner_id, created_at=stamp)
                    conn.execute(f'INSERT INTO partners ({",".join(fields)}) VALUES ({",".join("?" for _ in fields)})', tuple(fields.values()))
                    result['created'] += 1
            except ValidationError as exc:
                field = str(exc.errors()[0]['loc'][0]) if exc.errors()[0]['loc'] else 'name'
                fail(422, f'“伙伴信息”第{line}行：{COLUMNS.get(field, "内容")}无效，请检查必填项和标准值')
            except (HTTPException, ValueError) as exc:
                fail(exc.status_code if isinstance(exc, HTTPException) else 422,
                     f'“伙伴信息”第{line}行：{exc.detail if isinstance(exc, HTTPException) else "分类格式无效"}')
        if result['created'] or result['updated']:
            record_audit(conn, 'partners_imported', actor_user_id=actor, summary=result)
    return result

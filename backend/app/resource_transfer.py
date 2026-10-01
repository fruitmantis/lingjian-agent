"""Admin Excel maintenance of current course/lab fields; no database migration."""
from io import BytesIO
import json
import re
from typing import Literal
import uuid
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from pydantic import Field, ValidationError

from . import enablement as service, resource_categories as categories
from .database import get_db

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 2000
FLAGS = ('system_visible', 'model_allowed', 'partner_allowed')
# Column names are the file contract. A downloaded empty workbook is also a template.
COLUMNS = {
    'source_id': '资源ID（新增可留空）', 'resource_type': '类型', 'title': '名称',
    'summary': '简介', 'role_ids': '岗位（每行一个）', 'zone_ids': '专区（每行一个）',
    'level': '层级', 'duration_minutes': '实验时长（分钟）', 'source_url': '跳转链接',
    'course_goals': '课程目标', 'audience': '目标学员', 'outline': '课程大纲',
    'cover_url': '封面链接', 'lab_goals': '实验目标', 'lab_requirements': '基本要求',
    'system_visible': '系统内可见', 'model_allowed': '允许发送模型',
    'partner_allowed': '允许对伙伴外发', 'status': '导出时状态（仅供参考）',
}
STATUS = {'draft': '草稿', 'published': '已发布', 'unpublished': '已下架', 'revoked': '已撤销授权'}
KIND = {'course': '课程', 'lab': '实验'}
LEVEL = {'basic': '基础', 'advanced': '进阶', None: ''}
CATEGORY_HEADERS = ['类型', '名称', '排序']


def append_text(sheet, values):
    """Write untrusted text as strings, including strings beginning with '='."""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def export_workbook(*, template=False):
    book = Workbook()
    sheet = book.active
    sheet.title = '课程与实验'
    append_text(sheet, list(COLUMNS.values()))
    with get_db() as conn:
        conn.begin_read()
        taxonomy = categories.read(conn)
        names = {r['id']: r['name'] for r in taxonomy}
        for row in (() if template else conn.execute('SELECT * FROM enablement_resources ORDER BY created_at,id')):
            data = service.resource_metadata(json.loads(row['draft_json']))
            data.update(source_id=row['id'], status=STATUS[row['status']])
            for flag in FLAGS:
                data[flag] = '是' if row[flag] else '否'
            for field in ('role_ids', 'zone_ids'):
                if any(value not in names for value in data[field]):
                    service.fail(409, '资源存在失效分类，请先修正分类后导出')
                data[field] = '\n'.join(names[value] for value in data[field])
            data['resource_type'] = KIND[data['resource_type']]
            data['level'] = LEVEL[data['level']]
            if data['resource_type'] == '课程':
                data['duration_minutes'] = None
            append_text(sheet, [data.get(key, '') for key in COLUMNS])
    category_sheet = book.create_sheet('岗位与专区')
    append_text(category_sheet, CATEGORY_HEADERS)
    for row in taxonomy:
        append_text(category_sheet, ['岗位' if row['kind'] == 'role' else '专区', row['name'], row['sort_order']])
    notes = book.create_sheet('使用说明')
    for line in [
        '导出包含全部课程与实验的当前编辑资料、岗位/专区分类和三个独立用途授权。',
        '保留表头；类型填写课程或实验，层级填写基础或进阶，授权填写是或否。岗位/专区多选时每行一个名称。',
        '导入保存草稿，不自动上架；已发布资源的内容仍使用原发布版本，选择批量上架后才发布保存的草稿。',
        '授权按表格更新，减少授权会立即限制旧版本使用。批量上架不会自动打开用途授权。',
        '资源ID相同则更新；ID留空时按类型和跳转链接匹配，匹配不到才新增。不要更改已有资源ID。',
        '分类按类型和名称复用，缺失分类自动补充；已有分类的名称和排序保留。每个分类名称占一行。',
        '导出时状态仅供参考，不用于自动上架、下架或撤权。历史发布版本、审计和任务引用留在各自环境，不随表格迁移。',
        '任一行无效则整批不导入。只支持.xlsx，每次最多2000条、10MB；不支持公式。',
        '在内网部署包含本功能的同版本代码后导入；跳转链接和封面链接按原值保留，不下载来源平台的文件。',
    ]:
        append_text(notes, [line])
    for ws in book:
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.font = Font(bold=True, color='C7000B')
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 26
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        if ws != notes:
            ws.auto_filter.ref = ws.dimensions
    notes.column_dimensions['A'].width = 110
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def read_sheet(book, name, headers, limit):
    if name not in book.sheetnames:
        service.fail(422, f'缺少“{name}”工作表，请使用导出的模板')
    sheet = book[name]
    # Do not trust caller-controlled worksheet dimensions.
    sheet.reset_dimensions()
    rows = sheet.iter_rows()
    first = next(rows, ())
    if [c.value for c in first] != headers:
        service.fail(422, f'“{name}”表头不匹配，请使用导出的模板')
    result = []
    for line, cells in enumerate(rows, 2):
        if line > limit + 1:
            service.fail(422, f'“{name}”最多支持{limit}行')
        if any(c.data_type in ('f', 'e') for c in cells):
            service.fail(422, f'“{name}”第{line}行：请使用普通值，不要使用公式或错误值')
        values = [c.value for c in cells]
        if not any(value is not None for value in values):
            continue
        if len(values) > len(headers) and any(value is not None for value in values[len(headers):]):
            service.fail(422, f'“{name}”第{line}行：存在未知列')
        values = (values + [None] * len(headers))[:len(headers)]
        if any(value is not None and not isinstance(value, (str, int, float, bool)) for value in values):
            service.fail(422, f'“{name}”第{line}行：单元格格式无效')
        result.append((line, values))
    return result


def text(value):
    return '' if value is None else str(value).strip()


def read_workbook(content, sheets):
    if len(content) > MAX_BYTES:
        service.fail(413, '文件不能超过10MB')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 100 or sum(f.file_size for f in archive.infolist()) > 50 * 1024 * 1024:
                service.fail(422, '文件展开后过大，请拆分导入')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            result = {name: read_sheet(book, name, headers, limit) for name, headers, limit in sheets}
        finally:
            book.close()
    except HTTPException:
        raise
    except Exception:
        # Third-party XML parsers use several exception classes. Never expose their messages.
        service.fail(422, '无法读取Excel文件，请上传导出的.xlsx模板')
    return result


def parse_workbook(content):
    sheets = read_workbook(content, [('课程与实验', list(COLUMNS.values()), MAX_ROWS), ('岗位与专区', CATEGORY_HEADERS, 500)])
    return sheets['课程与实验'], sheets['岗位与专区']


def import_workbook(content, actor):
    resource_rows, category_rows = parse_workbook(content)
    if not resource_rows:
        service.fail(422, '表格中没有课程或实验')
    result = {'created': 0, 'updated': 0, 'unchanged': 0}
    with service.write_transaction() as conn:
        taxonomy = categories.read(conn)
        by_name = {(r['kind'], r['name'].casefold()): r for r in taxonomy}
        added = 0
        seen_categories = set()
        for line, (kind, name, order) in category_rows:
            try:
                item = categories.CategoryCreate(kind={'岗位': 'role', '专区': 'zone'}.get(text(kind), ''), name=text(name), sort_order=0 if order is None else order)
            except ValidationError:
                service.fail(422, f'“岗位与专区”第{line}行：请检查类型、名称和排序')
            key = (item.kind, item.name.casefold())
            if key in seen_categories:
                service.fail(422, f'“岗位与专区”第{line}行：分类重复')
            seen_categories.add(key)
            if key not in by_name:
                new = {'id': str(uuid.uuid4()), **item.model_dump()}
                taxonomy.append(new)
                by_name[key] = new
                added += 1
        if len(taxonomy) > 500:
            service.fail(422, '导入后分类数量超过500个，请合并重复分类')
        if added:
            conn.execute('INSERT INTO app_metadata (key,value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',
                         (categories.KEY, json.dumps(taxonomy, ensure_ascii=False)))
        seen_ids = set()
        for line, values in resource_rows:
            try:
                data = dict(zip(COLUMNS, values))
                source_id = text(data.pop('source_id'))
                if source_id and not re.fullmatch(r'[A-Za-z0-9_-]{1,100}', source_id):
                    service.fail(422, '资源ID无效')
                status = data.pop('status')
                if text(status) not in ('', *STATUS.values()):
                    service.fail(422, '导出时状态无效')
                flags = {}
                for flag in FLAGS:
                    value = text(data.pop(flag))
                    if value not in ('是', '否'):
                        service.fail(422, f'{COLUMNS[flag]}须填写是或否')
                    flags[flag] = value == '是'
                data['resource_type'] = {'课程': 'course', '实验': 'lab'}.get(text(data['resource_type']), '')
                data['level'] = {'基础': 'basic', '进阶': 'advanced', '': None}.get(text(data['level']), 'invalid')
                for field, kind in [('role_ids', 'role'), ('zone_ids', 'zone')]:
                    names = [name.strip() for name in text(data[field]).splitlines() if name.strip()]
                    if any((kind, name.casefold()) not in by_name for name in names):
                        service.fail(422, '岗位或专区不存在，请同时填写“岗位与专区”工作表')
                    data[field] = [by_name[(kind, name.casefold())]['id'] for name in names]
                for field in ('title', 'summary', 'source_url', 'course_goals', 'audience', 'outline', 'cover_url', 'lab_goals', 'lab_requirements'):
                    data[field] = text(data[field])
                if data['duration_minutes'] == '':
                    data['duration_minutes'] = None
                if data['resource_type'] == 'course' and data['duration_minutes'] is not None:
                    service.fail(422, '课程不填写实验时长')
                metadata = service.ResourceMetadata.model_validate(data)
                # Blank IDs support hand-authored rows and safe retries of the same file.
                if not source_id:
                    matches = []
                    for candidate in conn.execute('SELECT id,draft_json FROM enablement_resources'):
                        current = json.loads(candidate['draft_json'])
                        if current['resource_type'] == metadata.resource_type and current['source_url'] == metadata.source_url:
                            matches.append(candidate['id'])
                    if len(matches) > 1:
                        service.fail(409, '相同类型和链接对应多条资源，请填写资源ID')
                    source_id = matches[0] if matches else str(uuid.uuid4())
                if source_id in seen_ids:
                    service.fail(422, '资源重复，请每条资源只保留一行')
                seen_ids.add(source_id)
                existing = conn.execute('SELECT * FROM enablement_resources WHERE id=?', (source_id,)).fetchone()
                unchanged = existing and service.resource_metadata(json.loads(existing['draft_json'])) == metadata.model_dump() and all(bool(existing[f]) == flags[f] for f in FLAGS)
                if unchanged:
                    result['unchanged'] += 1
                    continue
                row = service.save('resource', source_id, service.ResourceSave(base_revision=existing['revision'] if existing else 0, metadata=metadata), actor, connection=conn)
                if any(bool(row[f]) != flags[f] for f in FLAGS):
                    service.permissions('resource', source_id, service.Permissions(base_revision=row['revision'], reason='Excel导入用途授权', **flags), actor, connection=conn)
                result['updated' if existing else 'created'] += 1
            except ValidationError as exc:
                field = str(exc.errors()[0]['loc'][0]) if exc.errors()[0]['loc'] else ''
                service.fail(422, f'“课程与实验”第{line}行：{COLUMNS.get(field, "资源信息")}无效，请检查内容、长度或链接')
            except HTTPException as exc:
                service.fail(exc.status_code, f'“课程与实验”第{line}行：{exc.detail}')
    return {**result, 'categories_added': added}


class BatchItem(service.Revision):
    source_id: str = Field(min_length=1, max_length=100)


class BatchAction(service.StrictModel):
    action: Literal['publish', 'unpublish']
    items: list[BatchItem] = Field(min_length=1, max_length=MAX_ROWS)


def batch_action(payload, actor):
    ids = [item.source_id for item in payload.items]
    if len(ids) != len(set(ids)):
        service.fail(422, '不能重复选择资源')
    changed = 0
    with service.write_transaction() as conn:
        for item in payload.items:
            row = service.row_for(conn, 'resource', item.source_id)
            service.check_base(row, item.base_revision)
        for item in payload.items:
            row = service.row_for(conn, 'resource', item.source_id)
            try:
                if payload.action == 'publish':
                    service.publish('resource', item.source_id, service.Revision(base_revision=item.base_revision), actor, connection=conn)
                elif row['status'] == 'published':
                    service.unpublish('resource', item.source_id, service.Unpublish(base_revision=item.base_revision, reason='管理员批量下架'), actor, connection=conn)
                else:
                    continue
                changed += 1
            except HTTPException as exc:
                title = service.resource_metadata(json.loads(row['draft_json']))['title']
                service.fail(exc.status_code, f'“{title}”：{exc.detail}；本批未生效')
    return {'changed': changed, 'skipped': len(ids) - changed}

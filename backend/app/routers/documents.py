"""Partner document upload and management router - requires authentication."""

import uuid
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, HTTPException, UploadFile, File, Depends, status
from fastapi.responses import FileResponse

from ..auth import require_auth
from ..database import get_db, UPLOADS_DIR
from ..doc_extractor import extract_text, get_file_type
from ..models import PartnerDocumentOut

router = APIRouter(prefix="/partners", tags=["documents"], dependencies=[Depends(require_auth)])
_DOC_COLS = "id, partner_id, filename, file_path, file_type, doc_category, extracted_text, created_at"
_ALLOWED_TYPES = {"pdf", "docx", "pptx", "xlsx"}


@router.get("/{partner_id}/documents", response_model=list[PartnerDocumentOut])
def list_documents(partner_id: str) -> list[PartnerDocumentOut]:
    with get_db() as conn:
        rows = conn.execute("SELECT id, partner_id, filename, file_type, doc_category, extracted_text, created_at FROM partner_documents WHERE partner_id = ? ORDER BY created_at DESC", (partner_id,)).fetchall()
    return [PartnerDocumentOut(**dict(r)) for r in rows]


@router.post("/{partner_id}/documents", response_model=PartnerDocumentOut, status_code=status.HTTP_201_CREATED)
async def upload_document(partner_id: str, file: UploadFile = File(...)) -> PartnerDocumentOut:
    if not file.filename:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="No filename")
    file_type = get_file_type(file.filename)
    if file_type not in _ALLOWED_TYPES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=f"Unsupported file type. Allowed: {', '.join(_ALLOWED_TYPES)}")
    with get_db() as conn:
        partner = conn.execute("SELECT id FROM partners WHERE id = ?", (partner_id,)).fetchone()
        if partner is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Partner not found")
    doc_id = str(uuid.uuid4())
    safe_name = Path(file.filename).name
    stored_name = f"{doc_id}_{safe_name}"
    file_path = UPLOADS_DIR / stored_name
    content = await file.read()
    file_path.write_bytes(content)
    extracted = extract_text(str(file_path), file_type)
    doc = PartnerDocumentOut(id=doc_id, partner_id=partner_id, filename=file.filename, file_type=file_type, doc_category=None, extracted_text=extracted, created_at=datetime.now(timezone.utc).isoformat())
    with get_db() as conn:
        conn.execute(f"INSERT INTO partner_documents ({_DOC_COLS}) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (doc.id, doc.partner_id, doc.filename, str(file_path), doc.file_type, doc.doc_category, doc.extracted_text, doc.created_at))
    return doc


@router.get("/{partner_id}/documents/{doc_id}/preview")
def preview_document(partner_id: str, doc_id: str):
    """Return HTML preview of a document (PPTX supported via python-pptx)."""
    with get_db() as conn:
        row = conn.execute("SELECT file_path, filename, file_type FROM partner_documents WHERE id = ? AND partner_id = ?", (doc_id, partner_id)).fetchone()
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Document not found")
    file_path = Path(row["file_path"])
    if not file_path.exists():
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="File not found on disk")

    if row["file_type"] == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation(str(file_path))
        slides_html = []
        for i, slide in enumerate(prs.slides, 1):
            shapes_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            shapes_text.append(f"<p>{text}</p>")
                elif hasattr(shape, "text") and shape.text.strip():
                    shapes_text.append(f"<p>{shape.text.strip()}</p>")
            slides_html.append(
                f'<div style="border:1px solid #ddd;border-radius:8px;padding:20px;margin-bottom:16px;min-height:200px;background:white">'
                f'<div style="font-size:12px;color:#999;margin-bottom:8px">幻灯片 {i}</div>'
                f'{"".join(shapes_text) if shapes_text else "<p style=\"color:#ccc\">空白幻灯片</p>"}'
                f'</div>'
            )
        html = f'<div style="font-family:sans-serif">{" ".join(slides_html)}</div>'
        from fastapi.responses import HTMLResponse
        return HTMLResponse(content=html)
    raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Preview not supported for this file type")


@router.get("/{partner_id}/documents/{doc_id}/file")
def download_document(partner_id: str, doc_id: str):
    with get_db() as conn:
        row = conn.execute("SELECT file_path, filename, file_type FROM partner_documents WHERE id = ? AND partner_id = ?", (doc_id, partner_id)).fetchone()
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Document not found")
    file_path = Path(row["file_path"])
    if not file_path.exists():
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="File not found on disk")
    return FileResponse(path=str(file_path), filename=row["filename"], media_type="application/octet-stream")


@router.delete("/{partner_id}/documents/{doc_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_document(partner_id: str, doc_id: str):
    with get_db() as conn:
        row = conn.execute("SELECT file_path FROM partner_documents WHERE id = ? AND partner_id = ?", (doc_id, partner_id)).fetchone()
        if row is None:
            raise HTTPException(status.HTTP_404_NOT_FOUND, detail="Document not found")
        try:
            Path(row["file_path"]).unlink(missing_ok=True)
        except Exception:
            pass
        conn.execute("DELETE FROM partner_documents WHERE id = ? AND partner_id = ?", (doc_id, partner_id))
    return None

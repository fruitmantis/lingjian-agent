"""Manual short-profile contracts on isolated PostgreSQL, with synthetic model output."""
import json
import os
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import threading
import pytest
from backend.app import profile_sources as sources,profile_report
from backend.app.database import get_db
from .test_profile_report import setup,upload,profile,organize,merge_result


@pytest.fixture
def short(setup,monkeypatch):
    calls=[];outputs={}
    def complete(_config,messages,_schema):
        data=json.loads(messages[-1]['content'])
        stage='source' if 'material' in data else 'merge'
        calls.append((stage,data))
        if stage=='merge':return json.dumps(merge_result(data),ensure_ascii=False)
        material=data['material']
        sections=outputs.get(material,[{'chapter':5,'summary':'内部工作台整合流程可视化、系统集成与AI分析。','quotes':[material]}])
        return json.dumps({'sections':sections},ensure_ascii=False)
    monkeypatch.setattr(sources.development_model,'completion',complete)
    return setup,calls,outputs,complete


def stored(pid):
    with get_db() as conn:
        return dict(conn.execute('SELECT * FROM partners WHERE id=?',(pid,)).fetchone())


def test_ten_sources_extract_independently_then_one_manual_merge_without_growth(short):
    ctx,calls,_,_=short;c,a,_,pid,_=ctx
    ids=[upload(ctx,f'flow-{n}.txt',f'内部演示材料{n}。流程图用于编排，连接系统，并用AI分析。') for n in range(10)]
    assert len(calls)==10 and all(stage=='source' for stage,_ in calls)
    assert stored(pid)['ai_profile'] is None and profile(ctx)['profile_needs_update']
    with get_db() as conn:
        for fid in ids:
            current=sources.source(conn,pid,'document',fid)
            sources.process_source(pid,'document',fid,sources.stamp(current))
    assert len(calls)==10
    assert c.post('/partners/batch-profile',headers=a).status_code==200
    assert len(calls)==10 and stored(pid)['ai_profile'] is None
    organize(ctx)
    assert len(calls)==11 and calls[-1][0]=='merge'
    body=profile(ctx)['ai_profile']
    assert body.count('内部工作台整合流程可视化、系统集成与AI分析。')==1
    assert all('材料'+str(n) not in body for n in range(10))
    assert not profile(ctx)['profile_needs_update']
    listed=c.get('/partners',headers=a)
    assert listed.status_code==200
    assert not next(row for row in listed.json() if row['id']==pid)['profile_needs_update']
    evidence=os.getenv('BANFEI_PROFILE_SHORT_EVIDENCE_DIR')
    if evidence:
        Path(evidence,'synthetic-saved-report.json').write_text(json.dumps(profile(ctx),ensure_ascii=False,indent=2))
    organize(ctx)
    assert len(calls)==11 and profile(ctx)['ai_profile']==body
    payload=calls[-1][1]
    assert 'material' not in payload and 'current_report' not in payload and 'intro' not in payload
    assert all(set(item)=={'kind','id','version','summary'} for chapter in payload['chapters'] for item in chapter['sources'])


def test_summary_is_not_a_continuous_quote_or_a_matching_input(short):
    ctx,calls,outputs,_=short;pid=ctx[3]
    original='内部演示。OriginalRareFalcon流程图用于编排，连接系统，用AI分析。'
    summary='内部演示工作台支持统一式数据流转与智能分析。'
    outputs[original]=[{'chapter':1,'summary':summary,'quotes':[original]}]
    fid=upload(ctx,'native.txt',original)
    assert summary not in original and stored(pid)['ai_profile'] is None
    with get_db() as conn:
        row=conn.execute('SELECT sections_json FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone()
        section=json.loads(row['sections_json'])[0]
        assert section['summary']==summary and all(q in original for q in section['quotes'])
        values=sources.matching_values(conn,pid)
        assert [v['path'] for v in values]==['/0/quotes/0']
        assert values[0]['text']==original and values[0]['original']==original
        assert values[0]['source_version']==sources.stamp(sources.source(conn,pid,'document',fid))
        assert not any(summary in value['text'] for value in values)
    organize(ctx)
    assert profile(ctx)['intro']==summary
    assert profile_report.parse(profile(ctx)['ai_profile']).chapters[0].split('\n',1)[1].strip()==summary
    assert calls[-1][1]['chapters'][0]['sources'][0]['summary']==summary
    assert original not in json.dumps(calls[-1][1],ensure_ascii=False)


def test_manual_failure_preserves_valid_body_and_ready_sources(short,monkeypatch):
    ctx,calls,_,complete=short;pid=ctx[3]
    upload(ctx,'first.txt','原有内部工作台资料。');organize(ctx)
    before=stored(pid)
    fid=upload(ctx,'new.txt','新增内部实践资料。')
    assert stored(pid)['ai_profile']==before['ai_profile']
    def fail(_config,messages,_schema):
        assert 'material' not in json.loads(messages[-1]['content'])
        raise TimeoutError('synthetic merge-only failure')
    monkeypatch.setattr(sources.development_model,'completion',fail)
    response=ctx[0].post(f'/partners/{pid}/profile',headers=ctx[1])
    assert response.status_code==502
    after=stored(pid)
    assert after['ai_profile']==before['ai_profile'] and after['profile_chapter_meta']==before['profile_chapter_meta']
    with get_db() as conn:
        assert conn.execute('SELECT state FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone()[0]=='ready'
    count=len([entry for entry in calls if entry[0]=='source'])
    monkeypatch.setattr(sources.development_model,'completion',complete)
    organize(ctx)
    assert len([entry for entry in calls if entry[0]=='source'])==count
    assert not profile(ctx)['profile_needs_update']


def test_delete_shared_source_invalidates_only_its_chapter_then_manual_rewrites(short):
    ctx,calls,outputs,_=short;c,a,_,pid,_=ctx
    company='公司在本地提供平台工程服务。'
    outputs[company]=[{'chapter':1,'summary':'本地平台工程伙伴。','quotes':[company]}]
    upload(ctx,'company.txt',company)
    first=upload(ctx,'flow-a.txt','共享流程能力资料甲。')
    second=upload(ctx,'flow-b.txt','共享流程能力资料乙。')
    organize(ctx);before=profile_report.parse(profile(ctx)['ai_profile']).chapters
    count=len(calls)
    assert c.delete(f'/partners/{pid}/documents/{first}',headers=a).status_code==204
    changed=profile_report.parse(profile(ctx)['ai_profile']).chapters
    assert changed[0]==before[0] and sources.CHANGED in changed[4]
    assert len(calls)==count and profile(ctx)['profile_needs_update']
    organize(ctx)
    assert profile_report.parse(profile(ctx)['ai_profile']).chapters[0]==before[0]
    assert '内部工作台' in profile(ctx)['ai_profile']
    count=len(calls)
    assert c.delete(f'/partners/{pid}/documents/{second}',headers=a).status_code==204
    final=profile_report.parse(profile(ctx)['ai_profile']).chapters
    assert final[0]==before[0] and final[4].split('\n',1)[1].strip()==profile_report.MISSING
    organize(ctx)
    assert len(calls)==count and '内部工作台' not in profile(ctx)['ai_profile']


@pytest.mark.parametrize('action',['replace','delete'])
def test_late_merge_commits_only_unchanged_chapters_without_an_automatic_loop(short,monkeypatch,action):
    ctx,calls,outputs,complete=short;c,a,_,pid,_=ctx
    company='公司提供本地工程服务。'
    outputs[company]=[{'chapter':1,'summary':'本地工程服务伙伴。','quotes':[company]}]
    upload(ctx,'company.txt',company)
    fid=upload(ctx,'old.txt','OldNativeFlow资料。')
    entered=threading.Event();release=threading.Event();merges=[]
    def waiting(_config,messages,schema):
        data=json.loads(messages[-1]['content'])
        if 'material' in data:return complete(_config,messages,schema)
        merges.append(data);entered.set();assert release.wait(20)
        return json.dumps(merge_result(data))
    monkeypatch.setattr(sources.development_model,'completion',waiting)
    with ThreadPoolExecutor(max_workers=1) as pool:
        request=pool.submit(c.post,f'/partners/{pid}/profile',headers=a)
        try:
            assert entered.wait(10)
            if action=='replace':
                replacement='NewNativeFlow仅内部演示。'
                outputs[replacement]=[{'chapter':5,'summary':'新的内部演示流程。','quotes':[replacement]}]
                response=c.put(f'/partners/{pid}/documents/{fid}',headers=a,files={'file':('new.txt',replacement.encode())})
                assert response.status_code==200
            else:assert c.delete(f'/partners/{pid}/documents/{fid}',headers=a).status_code==204
        finally:release.set()
        assert request.result(timeout=10).status_code==200
    assert len(merges)==1
    body=profile(ctx)['ai_profile']
    assert '本地工程服务伙伴' in body and '内部工作台' not in body
    if action=='replace':
        assert profile(ctx)['profile_needs_update']
        monkeypatch.setattr(sources.development_model,'completion',complete)
        organize(ctx)
        assert '新的内部演示流程' in profile(ctx)['ai_profile']
    assert 'OldNativeFlow' not in profile(ctx)['ai_profile']


def test_newer_same_input_manual_result_cannot_be_overwritten(short,monkeypatch):
    ctx,_,_,complete=short;c,a,_,pid,_=ctx
    upload(ctx,'flow.txt','当前内部流程资料。')
    replies=[]
    def competing(_config,messages,schema):
        data=json.loads(messages[-1]['content'])
        if 'material' in data:return complete(_config,messages,schema)
        replies.append(data)
        result=merge_result(data)
        if len(replies)==1:
            assert c.post(f'/partners/{pid}/profile',headers=a).status_code==200
            result['sections'][0]['body']='旧请求的内部流程结果。'
        else:result['sections'][0]['body']='新请求的内部流程结果。'
        return json.dumps(result)
    monkeypatch.setattr(sources.development_model,'completion',competing)
    assert c.post(f'/partners/{pid}/profile',headers=a).status_code==200
    assert len(replies)==2
    assert '新请求的内部流程结果' in profile(ctx)['ai_profile']
    assert '旧请求的内部流程结果' not in profile(ctx)['ai_profile']
    assert not profile(ctx)['profile_needs_update']


def test_legacy_backfill_is_selected_and_keeps_original_quotes_order_and_state(short):
    ctx,calls,outputs,_=short;pid=ctx[3]
    original='第一项内部事实。\n\n第二项演示事实。'
    fid=upload(ctx,'legacy.txt',original)
    old=[{'chapter':5,'quotes':['第二项演示事实。','第一项内部事实。']}]
    with get_db() as conn:
        current=sources.source(conn,pid,'document',fid);version=sources.stamp(current)
        sources.put(conn,current,'ready',old)
        before=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
    calls.clear()
    organize(ctx)
    assert calls==[] and stored(pid)['ai_profile'] is None
    sources.backfill_summary(pid,'document',fid,version)
    assert len(calls)==1 and calls[0][0]=='source'
    with get_db() as conn:
        after=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
        updated=json.loads(after['sections_json'])
        assert updated[0]['quotes']==old[0]['quotes']
        assert (after['state'],after['source_fingerprint'],after['error'])==(before['state'],before['source_fingerprint'],before['error'])
        assert [r['path'] for r in sources.matching_values(conn,pid)]==['/0/quotes/0','/0/quotes/1']
    assert stored(pid)['ai_profile'] is None
    sources.backfill_summary(pid,'document',fid,version)
    assert len(calls)==1
    organize(ctx)
    assert len(calls)==2 and calls[-1][0]=='merge'


def test_invalid_or_late_legacy_backfill_never_resets_quotes_or_state(short,monkeypatch):
    ctx,_,_,_=short;pid=ctx[3]
    fid=upload(ctx,'legacy.txt','现有内部事实。')
    with get_db() as conn:
        current=sources.source(conn,pid,'document',fid);version=sources.stamp(current)
        sources.put(conn,current,'ready',[{'chapter':5,'quotes':['现有内部事实。']}],error='existing partial-quote note')
        before=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
    monkeypatch.setattr(sources.development_model,'completion',lambda *_:json.dumps({'sections':[{'chapter':5,'summary':'不成立的摘要','quotes':['不存在的引用。']}]}))
    with pytest.raises(ValueError):sources.backfill_summary(pid,'document',fid,version)
    with get_db() as conn:
        assert dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())==before
    def changed(*_):
        with get_db() as conn:
            conn.execute('UPDATE partner_documents SET extracted_text=?,file_path=? WHERE id=?',('替换后的原文。','synthetic-new-version',fid))
            from backend.app.material_files import changed as mark
            mark(conn,pid)
        return json.dumps({'sections':[{'chapter':5,'summary':'旧事实摘要','quotes':['现有内部事实。']}]})
    monkeypatch.setattr(sources.development_model,'completion',changed)
    sources.backfill_summary(pid,'document',fid,version)
    with get_db() as conn:
        now=conn.execute('SELECT source_fingerprint,state,sections_json FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone()
        assert now['source_fingerprint']!=version and now['state']=='pending'
        assert now['sections_json']=='[]'


def test_sources_mark_only_actual_used_refs_and_recheck_visibility(short,monkeypatch):
    ctx,_,_,complete=short;c,a,u,pid,_=ctx
    ids=[]
    for title,visible in [('PUBLIC_USED_CASE',True),('HIDDEN_UNUSED_CASE',False)]:
        case=c.post('/cases',headers=a,json={'partner_id':pid,'title':title,'category_id':'technical-1','visible':visible})
        cid=case.json()['id']
        response=c.post(f'/cases/{cid}/deliverables',headers=a,files={'file':(title+'.txt',(title+'内部资料。').encode())})
        ids.append((cid,response.json()['id']))
    def selected(_config,messages,schema):
        data=json.loads(messages[-1]['content'])
        if 'material' in data:return complete(_config,messages,schema)
        entry=next(item for item in data['chapters'][0]['sources'] if item['id']==ids[0][1])
        return json.dumps({'sections':[{'chapter':5,'body':entry['summary'],'used_sources':[{key:entry[key] for key in ('kind','id','version')}]}]})
    monkeypatch.setattr(sources.development_model,'completion',selected)
    organize(ctx)
    meta=json.loads(stored(pid)['profile_chapter_meta'])
    assert [r['id'] for r in meta['5']['used_sources']]==[ids[0][1]]
    admin=profile(ctx,True)['ai_profile']
    assert ids[0][1] in admin.split('已参与：')[1].split('未参与：')[0]
    assert ids[1][1] in admin.split('未参与：')[1]
    public=profile(ctx)
    assert 'HIDDEN_UNUSED_CASE' not in json.dumps(public)
    assert 'profile_chapter_meta' not in public and public.get('profile_sources') is None
    assert c.patch(f'/cases/{ids[0][0]}/visibility',headers=a,json={'visible':False}).status_code==200
    assert 'PUBLIC_USED_CASE' not in json.dumps(profile(ctx))
    assert '内部工作台' in profile(ctx)['ai_profile']


def test_field_only_migration_preserves_existing_profiles_quotes_and_rows(short):
    from backend.app.profile_chapter_schema import migrate
    ctx,_,_,_=short;pid=ctx[3]
    fid=upload(ctx,'migration.txt','迁移前原文事实。')
    with get_db() as conn:
        conn.execute('ALTER TABLE partners DROP COLUMN profile_chapter_meta')
        before=dict(conn.execute('SELECT * FROM partners WHERE id=?',(pid,)).fetchone())
        source_before=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
        assert migrate(conn)['existing_profiles_and_sources_retained']
        after=dict(conn.execute('SELECT * FROM partners WHERE id=?',(pid,)).fetchone())
        assert json.loads(after.pop('profile_chapter_meta'))=={} and after==before
        assert dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())==source_before
        assert migrate(conn)['schema_version']==20
        assert conn.execute('SELECT count(*) FROM partners').fetchone()[0]==1


def test_short_bodies_do_not_require_empty_tables_and_keep_format_validation():
    for number in (1,4,6):profile_report.validate_body(number,'一条具体事实。')
    for bad in ('','{"raw":"json"}','## 5. AI技术能力与解决方案'):
        with pytest.raises(profile_report.ReportError):profile_report.validate_body(1,bad)
    text=profile_report.empty_report().text
    assert '|' not in text and '免责声明' not in text
    assert text.count(profile_report.MISSING)==10


def test_legacy_summary_stays_pending_after_another_chapter_and_noop_actions(short):
    ctx,calls,outputs,_=short;c,a,_,pid,_=ctx
    fid=upload(ctx,'legacy.txt','旧来源尚无摘要。')
    with get_db() as conn:
        source=sources.source(conn,pid,'document',fid)
        sources.put(conn,source,'ready',[{'chapter':5,'quotes':['旧来源尚无摘要。']}])
        before=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
    company='新增公司事实。'
    outputs[company]=[{'chapter':1,'summary':'本地工程伙伴。','quotes':[company]}]
    upload(ctx,'company.txt',company);organize(ctx)
    assert [stage for stage,_ in calls]==['source','source','merge']
    assert profile(ctx)['profile_needs_update'] and profile(ctx,True)['profile_needs_update']
    assert stored(pid)['profile_materials_revision']==-1
    listed=c.get('/partners',headers=a).json()
    assert next(row for row in listed if row['id']==pid)['profile_needs_update']
    materials=c.get('/admin/partner-materials',headers=a)
    assert materials.status_code==200
    assert next(row for row in materials.json()['items'] if row['id']==fid)['profile_needs_update']
    organize(ctx)
    assert c.post('/partners/batch-profile',headers=a).status_code==200
    assert len(calls)==3 and profile(ctx)['profile_needs_update']
    with get_db() as conn:
        assert dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())==before


@pytest.mark.parametrize('location',['prefix','chapter','section'])
def test_word_adoption_and_saved_projection_preserve_native_text_and_sources(short,location):
    import io
    from docx import Document
    ctx,calls,_,_=short;pid=ctx[3]
    disclaimer='免责声明：仅供参考，不承担任何责任。'
    doc=Document();doc.add_paragraph('合成伙伴，2026年10月。')
    if location=='prefix':doc.add_paragraph(disclaimer)
    for number,title in enumerate(profile_report.CHAPTERS,1):
        doc.add_heading(f'{number}. {title}',1)
        doc.add_paragraph('内部演示事实，尚未交付。' if number==1 else '合成事实。')
        if number==1 and location=='chapter':doc.add_paragraph(disclaimer)
        if number==1 and location=='section':
            doc.add_heading('免责声明',3);doc.add_paragraph('仅供参考，不承担任何责任。')
    stream=io.BytesIO();doc.save(stream);original=stream.getvalue()
    fid=upload(ctx,'synthetic.docx',original,True)
    assert calls==[]
    with get_db() as conn:
        source=sources.source(conn,pid,'word',fid)
        assert '免责声明' in source['text'] and Path(source['file_path']).read_bytes()==original
        row=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
        assert '免责声明' in row['sections_json']
        original_values=sources.matching_values(conn,pid)
        assert any('免责声明' in value['text'] for value in original_values)
    saved=stored(pid)['ai_profile']
    assert '免责声明' in saved and '不承担任何责任' in saved
    assert '合成伙伴，2026年10月。' in saved
    assert profile(ctx)['intro']==profile_report.parse(saved).chapters[0].split('\n',1)[1].strip()
    if location!='prefix':assert '不承担任何责任' in profile(ctx)['intro']
    assert '免责声明' in profile(ctx,True)['ai_profile']
    organize(ctx)
    assert calls==[] and not profile(ctx)['profile_needs_update']
    # Historical saved output stays readable without cleaning or rewriting raw data.
    with get_db() as conn:
        conn.execute('UPDATE partners SET ai_profile=? WHERE id=?',(disclaimer+'\n\n'+saved,pid))
    assert profile(ctx)['ai_profile'].startswith(disclaimer+'\n\n')
    assert profile(ctx,True)['ai_profile'].startswith(disclaimer+'\n\n')
    assert stored(pid)['ai_profile']==disclaimer+'\n\n'+saved
    with get_db() as conn:
        assert dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())==row
        assert sources.matching_values(conn,pid)==original_values


def test_source_and_manual_prompts_use_direct_facts_preserving_business_names(short,monkeypatch):
    ctx,calls,outputs,complete=short
    original='本公司提供合同免责声明检测平台；本次材料为内部演示，尚未对外交付。'
    outputs[original]=[{'chapter':1,'summary':original,'quotes':[original]}]
    messages_seen=[]
    def capture(config,messages,schema):
        messages_seen.append(messages)
        return complete(config,messages,schema)
    monkeypatch.setattr(sources.development_model,'completion',capture)
    upload(ctx,'synthetic.txt',original);organize(ctx)
    assert [stage for stage,_ in calls]==['source','merge'] and len(messages_seen)==2
    for messages in messages_seen:
        prompt=messages[0]['content']
        assert '直接陈述有依据的事实' in prompt and '保留完整业务名称与实际限定' in prompt
        assert '不添加推责、过度保守或否定式兜底套话' in prompt
        assert '删除全部免责声明' not in prompt and '省略导航、采集过程、泛泛宣传和全部免责声明' not in prompt
        assert original in json.dumps(json.loads(messages[-1]['content']),ensure_ascii=False)
    assert original in stored(ctx[3])['ai_profile'] and profile(ctx)['intro']==original
    with get_db() as conn:
        values=sources.matching_values(conn,ctx[3])
        assert any(v['text']==original and v['original']==original for v in values)
    organize(ctx);assert len(calls)==2 and len(messages_seen)==2


@pytest.mark.parametrize('fact',[
    '本公司已交付合同免责声明检测平台，支持合同条款分类。',
    '本公司提供 Disclaimer 合同分析模块，已部署于客户内网。',
])
def test_word_business_disclaimer_terms_survive_saving_projection_and_raw_source(short,fact):
    import io
    from docx import Document
    ctx,calls,_,_=short;pid=ctx[3]
    doc=Document();doc.add_paragraph('合成公司，2026年。')
    doc.add_paragraph(fact)
    doc.add_paragraph('免责声明：仅供参考，不承担任何责任。')
    module='免责声明检测平台' if '免责声明' in fact else 'Disclaimer 合同分析模块'
    for number,title in enumerate(profile_report.CHAPTERS,1):
        doc.add_heading(f'{number}. {title}',1)
        doc.add_paragraph('内部演示，尚未交付。')
        if number==5:
            doc.add_heading(module,3);doc.add_paragraph(fact)
    stream=io.BytesIO();doc.save(stream);original=stream.getvalue()
    fid=upload(ctx,'business-fact.docx',original,True)
    assert calls==[]
    saved=stored(pid)['ai_profile']
    assert fact in saved and module in profile_report.parse(saved).chapters[4]
    assert '仅供参考，不承担任何责任。' in saved
    for admin in (False,True):
        assert fact in profile(ctx,admin)['ai_profile']
        assert '仅供参考，不承担任何责任。' in profile(ctx,admin)['ai_profile']
    with get_db() as conn:
        raw=sources.source(conn,pid,'word',fid)
        assert Path(raw['file_path']).read_bytes()==original and fact in raw['text']
        row=conn.execute('SELECT sections_json FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone()
        assert fact in row['sections_json'] and '仅供参考，不承担任何责任。' in row['sections_json']
        values=sources.matching_values(conn,pid)
        assert any(fact in entry['text'] and entry['source_version']==sources.stamp(raw) for entry in values)
    organize(ctx)
    assert calls==[] and stored(pid)['ai_profile']==saved


def test_legacy_reclassification_preserves_quote_paths_and_merges_only_actual_summaries(short):
    ctx,calls,outputs,_=short;c,a,_,pid,_=ctx
    company='公司在本地提供平台工程服务。'
    flow='内部演示流程可视化。'
    case='仅展示内部实践案例，尚未对外交付。'
    boundary='采集时间：2026-10-09；HTTP/TLS记录。'
    collaboration='可以共同开展内部实验。'
    communication='沟通时确认演示版本。'
    original='\n'.join([company,flow,case,boundary,collaboration,communication])
    fid=upload(ctx,'legacy-reclassified.txt',original)
    old=[{'chapter':5,'quotes':[flow]},{'chapter':7,'quotes':[boundary]},
         {'chapter':8,'quotes':[collaboration]},{'chapter':9,'quotes':[communication]}]
    with get_db() as conn:
        source=sources.source(conn,pid,'document',fid);version=sources.stamp(source)
        sources.put(conn,source,'ready',old)
        before=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
        values=sources.matching_values(conn,pid)
    calls.clear()
    assert profile(ctx)['profile_needs_update']
    outputs[original]=[
        {'chapter':1,'summary':company,'quotes':[company]},
        {'chapter':5,'summary':flow,'quotes':[flow]},
        {'chapter':6,'summary':case,'quotes':[case]},
        {'chapter':8,'summary':collaboration,'quotes':[collaboration]},
        {'chapter':9,'summary':communication,'quotes':[communication]}]
    sources.backfill_summary(pid,'document',fid,version)
    assert [stage for stage,_ in calls]==['source']
    with get_db() as conn:
        after=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
        sections=json.loads(after['sections_json'])
        assert [{k:v for k,v in section.items() if k!='summary_sections'} for section in sections]==old
        assert sources.matching_values(conn,pid)==values
        assert (after['state'],after['source_fingerprint'],after['error'])==(before['state'],before['source_fingerprint'],before['error'])
        inputs=sources.chapter_inputs(conn,pid)
        assert [int(chapter) for chapter,entries in inputs.items() if entries]==[1,5,6,8,9]
        assert inputs['7']==[]
        assert not sources._missing_summaries(sources._ready_sources(conn,pid,sources.sources(conn,pid))[0])
    assert profile(ctx)['profile_needs_update']
    sources.process_source(pid,'document',fid,version)
    sources.backfill_summary(pid,'document',fid,version)
    organize(ctx)
    assert [stage for stage,_ in calls]==['source','merge']
    saved=profile_report.parse(profile(ctx)['ai_profile'])
    assert saved.chapters[6].split('\n',1)[1].strip()==profile_report.MISSING
    assert profile(ctx)['intro']==company and case in saved.chapters[5]
    assert not profile(ctx)['profile_needs_update'] and stored(pid)['profile_materials_revision']>=0
    organize(ctx);assert c.post('/partners/batch-profile',headers=a).status_code==200
    assert len(calls)==2


@pytest.mark.parametrize('has_facts',[True,False])
def test_legacy_saved_result_reuse_issues_zero_model_calls_and_records_empty_result(short,has_facts):
    ctx,calls,_,_=short;pid=ctx[3]
    original='本公司提供内部演示平台。\n采集记录：2026-10-09。'
    fid=upload(ctx,'legacy-reuse.txt',original)
    old=[{'chapter':7,'quotes':['采集记录：2026-10-09。']}]
    with get_db() as conn:
        source=sources.source(conn,pid,'document',fid);version=sources.stamp(source)
        sources.put(conn,source,'ready',old)
        values=sources.matching_values(conn,pid)
    calls.clear()
    result={'sections':[{'chapter':5,'summary':'本公司提供内部演示平台。','quotes':['本公司提供内部演示平台。']}] if has_facts else []}
    sources.backfill_summary(pid,'document',fid,version,raw_result=json.dumps(result,ensure_ascii=False))
    assert calls==[]
    with get_db() as conn:
        sections=json.loads(conn.execute('SELECT sections_json FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone()[0])
        assert sections[0]['quotes']==old[0]['quotes'] and sources.matching_values(conn,pid)==values
        assert not sources._missing_summaries(sources._ready_sources(conn,pid,sources.sources(conn,pid))[0])
        assert bool(sources.chapter_inputs(conn,pid)['5'])==has_facts
    sources.process_source(pid,'document',fid,version)
    sources.backfill_summary(pid,'document',fid,version)
    assert calls==[]
    organize(ctx)
    assert [stage for stage,_ in calls]==(['merge'] if has_facts else [])
    assert not profile(ctx)['profile_needs_update']
    count=len(calls);organize(ctx);assert len(calls)==count


@pytest.mark.parametrize('invalid_result',[
    {'sections':[{'chapter':5,'summary':'无原文支撑。','quotes':['原文不存在。']}]},
    {'sections':[{'chapter':5,'summary':'重复一。','quotes':['现有内部事实。']},
                 {'chapter':5,'summary':'重复二。','quotes':['现有内部事实。']}]},
])
def test_reused_legacy_result_retains_existing_validation_and_failure_preserves_rows(short,invalid_result):
    ctx,calls,_,_=short;pid=ctx[3];fid=upload(ctx,'invalid-reuse.txt','现有内部事实。')
    with get_db() as conn:
        source=sources.source(conn,pid,'document',fid);version=sources.stamp(source)
        sources.put(conn,source,'ready',[{'chapter':5,'quotes':['现有内部事实。']}])
        before=dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())
    calls.clear()
    with pytest.raises(ValueError):
        sources.backfill_summary(pid,'document',fid,version,raw_result=json.dumps(invalid_result,ensure_ascii=False))
    with get_db() as conn:
        assert dict(conn.execute('SELECT * FROM partner_profile_sources WHERE source_id=?',(fid,)).fetchone())==before
    assert calls==[] and profile(ctx)['profile_needs_update']


@pytest.mark.parametrize('admin',[False,True])
def test_development_context_intro_uses_current_visible_first_chapter_without_extra_fields(short,admin):
    ctx,calls,_,_=short;c,a,u,pid,_=ctx
    upload(ctx,'private-company.txt','内部公司资料。')
    report=profile_report.empty_report()
    report.chapters[0]=profile_report.heading(1)+'\n\n本公司提供合同免责声明检测平台；依据 private-company.txt。\n\n'
    report.chapters[4]=profile_report.heading(5)+'\n\n这是第5章，不是摘要卡片内容。\n\n'
    with get_db() as conn:
        conn.execute('UPDATE partners SET ai_profile=?,intro=? WHERE id=?',(report.text,'旧简介不应被摘要卡片采用。',pid))
    count=len(calls);headers=a if admin else u
    expected=c.get(f'/partners/{pid}',headers=headers)
    assert expected.status_code==200
    response=c.get(f'/enablement/context?partner_id={pid}',headers=headers)
    assert response.status_code==200
    partner=response.json()['partner']
    assert partner['intro']==expected.json()['intro'] and '合同免责声明检测平台' in partner['intro']
    assert '这是第5章' not in partner['intro'] and '旧简介' not in partner['intro']
    assert not {'profile_chapter_meta','profile_sources','profile_needs_update'} & partner.keys()
    if admin:assert partner['ai_profile']==report.text
    else:
        assert partner['ai_profile'] is None and 'private-company.txt' not in partner['intro']
    assert stored(pid)['intro']=='旧简介不应被摘要卡片采用。' and stored(pid)['ai_profile']==report.text
    assert len(calls)==count

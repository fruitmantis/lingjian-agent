"""Synthetic ten-chapter Word and model output, with no application/database setup."""
import io
import json
from docx import Document
from backend.app import profile_report as report


def document(*,toc=False,invalid=False,extra='',empty_cases=False):
    doc=Document();doc.add_paragraph('合成伙伴报告');doc.add_paragraph('来源日期：2026-09-24；集团口径，企业自述，待核实')
    if toc:
        doc.add_heading('目录',1)
        for n,t in enumerate(report.CHAPTERS,1):doc.add_paragraph(f'{n}. {t}')
    for n,t in enumerate(report.CHAPTERS,1):
        if invalid and n==7:continue
        title='四、华为认证与资质核查（专题）' if n==4 else f'{n}. {t}'
        doc.add_heading(title,1);doc.add_paragraph(f'第{n}章原有完整说明；来源日期2026-09-24，集团口径，企业自述。')
        if n == 1 and extra: doc.add_paragraph(extra)
        if n in (1,4,6):
            rows=[['字段','事实'],['限定','本公司 | 未核验\n现有资料未提供']]
            if n==6 and empty_cases: rows=rows[:1]
            table=doc.add_table(rows=len(rows),cols=2)
            for i,values in enumerate(rows):
                for j,value in enumerate(values):table.cell(i,j).text=value
    result=io.BytesIO();doc.save(result);return result.getvalue()


def patch_all(text='新报告说明'):
    base=report.empty_report()
    return json.dumps({'sections':[{'chapter':i,'content':chapter.split('\n',1)[1].strip()+f'\n\n{text}'} for i,chapter in enumerate(base.chapters,1)]},ensure_ascii=False)



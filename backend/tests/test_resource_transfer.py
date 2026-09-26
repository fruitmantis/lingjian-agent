"""Excel and batch operations against an isolated PostgreSQL validation schema."""
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest

from backend.app import enablement as service, resource_transfer as transfer, resource_categories as categories
from backend.app.database import get_db
from backend.tests.conftest import auth_headers, make_user
from backend.tests.test_enablement import admin, metadata, create, grant, published, resolve

PATH = '/admin/enablement/resources'


def workbook(content=None, *, edits=(), extra=(), category_rows=()):
    book = load_workbook(BytesIO(content or transfer.export_workbook()))
    sheet = book['课程与实验']
    for row, key, value in edits:
        sheet.cell(row, list(transfer.COLUMNS).index(key) + 1, value)
    for row in extra:
        transfer.append_text(sheet, row)
    for row in category_rows:
        transfer.append_text(book['岗位与专区'], row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def upload(client, admin, content):
    return client.post(PATH + '/import', headers=auth_headers(admin), files={'file': ('resources.xlsx', content)})


def counts():
    with get_db() as conn:
        return tuple(conn.execute(f'SELECT count(*) FROM {table}').fetchone()[0] for table in ('enablement_resources', 'enablement_resource_versions', 'enablement_audit_events'))


@pytest.mark.parametrize('kind', ['course', 'lab'])
def test_all_current_fields_excel_round_trip_and_new_drafts(client, admin, metadata, kind):
    metadata.update(resource_type=kind, level='advanced', title='=完整文字标题', summary='第一行\n第二行', role_ids=['role-1', 'role-4'], zone_ids=['zone-1', 'zone-2'])
    if kind == 'course':
        metadata.update(course_goals='目标', audience='学员', outline='大纲\n' + '字' * 11000, cover_url='https://example.com/cover.png')
    else:
        metadata.update(lab_goals='实验目标', lab_requirements='实验要求', duration_minutes=45)
    row = published(grant(create(admin, metadata), admin, flags=(True, False, True)), admin)
    response = client.get(PATH + '/export', headers=auth_headers(admin))
    assert response.status_code == 200 and 'spreadsheetml' in response.headers['content-type']
    book = load_workbook(BytesIO(response.content))
    assert book['课程与实验']['C2'].data_type == 's'  # No Excel formula injection.
    body = workbook(response.content, edits=[(2, 'source_id', 'imported-copy')])
    result = upload(client, admin, body)
    assert result.status_code == 200 and result.json()['created'] == 1
    imported = service.detail('resource', 'imported-copy')
    assert imported['metadata'] == row['metadata']
    assert imported['status'] == 'draft' and imported['published_version'] is None
    assert [bool(imported[f]) for f in transfer.FLAGS] == [True, False, True]
    assert upload(client, admin, body).json()['unchanged'] == 1
    assert service.detail('resource', 'imported-copy')['revision'] == imported['revision']
    assert counts()[:2] == (2, 1)


def test_update_draft_and_permissions_preserves_published_snapshot(client, admin, metadata):
    row = published(grant(create(admin, metadata), admin), admin)
    before = counts()
    body = workbook(edits=[(2, 'title', '导入的新草稿')])
    assert upload(client, admin, body).json()['updated'] == 1
    updated = service.detail('resource', row['source_id'])
    assert updated['metadata']['title'] == '导入的新草稿'
    assert resolve(updated)['title'] == metadata['title']
    assert counts()[1] == before[1]
    body = workbook(edits=[(2, 'model_allowed', '否')])
    assert upload(client, admin, body).status_code == 200
    updated = service.detail('resource', row['source_id'])
    assert not updated['model_allowed'] and updated['partner_allowed'] and updated['system_visible']
    with pytest.raises(HTTPException) as error:
        resolve(updated, purpose='model')
    assert error.value.status_code == 409
    assert updated['published_metadata']['title'] == metadata['title']


def test_category_names_map_across_environments_and_do_not_reset_existing(client, admin, metadata):
    create(admin, metadata)
    body = transfer.export_workbook()
    with get_db() as conn:
        rows = categories.read(conn)
        for category in rows:
            category['id'] = 'target-' + category['id']
            category['sort_order'] += 10
        conn.execute('UPDATE app_metadata SET value=? WHERE key=?', (json.dumps(rows), categories.KEY))
    body = workbook(body, edits=[(2, 'source_id', 'new-environment-resource')], category_rows=[['专区', '新专区', 25]])
    result = upload(client, admin, body)
    assert result.status_code == 200 and result.json()['categories_added'] == 1
    assert service.detail('resource', 'new-environment-resource')['metadata']['role_ids'] == ['target-role-1', 'target-role-4']
    assert next(c for c in categories.listing() if c['id'] == 'target-role-1')['sort_order'] == 10
    assert len(categories.listing()) == 12


def test_import_error_rolls_back_rows_categories_permissions_and_audit(client, admin, metadata):
    create(admin, metadata)
    original = transfer.export_workbook()
    book = load_workbook(BytesIO(original))
    new_row = [cell.value for cell in book['课程与实验'][2]]
    new_row[0] = 'bad-row'
    new_row[list(transfer.COLUMNS).index('source_url')] = 'http://127.0.0.1/secret'
    before = counts()
    body = workbook(original, edits=[(2, 'title', '不能保留的修改'), (2, 'system_visible', '是')], extra=[new_row], category_rows=[['岗位', '不能保留的分类', 1]])
    response = upload(client, admin, body)
    assert response.status_code == 422 and '第3行' in response.json()['detail']
    assert counts() == before
    assert service.detail('resource', 'resource-test')['metadata']['title'] == metadata['title']
    assert len(categories.listing()) == 11


def test_blank_ids_match_type_and_url_and_concurrent_retries_do_not_duplicate(client, admin, metadata):
    create(admin, metadata)
    body = workbook(edits=[(2, 'source_id', ''), (2, 'source_url', 'https://example.com/new')])
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: transfer.import_workbook(body, admin['id']), range(2)))
    assert sorted(r['created'] for r in results) == [0, 1]
    assert counts()[0] == 2


@pytest.mark.parametrize('key,value', [('source_url', 'javascript:alert(1)'), ('source_id', '../../invalid'), ('model_allowed', ''), ('role_ids', '不存在的分类'), ('level', '高级'), ('resource_type', '视频'), ('duration_minutes', 20), ('title', '=1+1')])
def test_invalid_cells_are_reported_with_row_number(client, admin, metadata, key, value):
    create(admin, metadata)
    before = counts()
    response = upload(client, admin, workbook(edits=[(2, key, value)]))
    assert response.status_code == 422 and '第2行' in response.json()['detail']
    assert counts() == before


def test_duplicate_rows_and_bad_files_rejected(client, admin, metadata):
    create(admin, metadata)
    book = load_workbook(BytesIO(transfer.export_workbook()))
    duplicate = [c.value for c in book['课程与实验'][2]]
    assert upload(client, admin, workbook(extra=[duplicate])).status_code == 422
    assert upload(client, admin, b'not a workbook').status_code == 422
    assert upload(client, admin, b'x' * (transfer.MAX_BYTES + 1)).status_code == 413
    stream = BytesIO()
    with ZipFile(stream, 'w') as archive:
        for i in range(101):
            archive.writestr(str(i), '')
    assert upload(client, admin, stream.getvalue()).status_code == 422


def test_batch_lifecycle_versions_permissions_and_stale_conflicts(client, admin, metadata):
    first = grant(create(admin, metadata), admin, flags=(True, False, True))
    second = service.save('resource', 'second', service.ResourceSave(base_revision=0, metadata={**metadata, 'resource_type': 'lab', 'source_url': 'https://example.com/lab'}), admin['id'])
    def batch(rows, action='publish'):
        return client.post(PATH + '/batch', headers=auth_headers(admin), json={'action': action, 'items': [{'source_id': r['source_id'], 'base_revision': r['revision']} for r in rows]})
    assert batch([first, second]).json() == {'changed': 2, 'skipped': 0}
    new_first = service.detail('resource', first['source_id'])
    new_second = service.detail('resource', second['source_id'])
    assert new_first['status'] == 'published' and new_second['status'] == 'published'
    assert not new_first['model_allowed'] and not new_second['system_visible']
    assert counts()[1] == 2
    before = counts()
    assert batch([new_second, first]).status_code == 409
    assert counts() == before
    assert batch([new_first, new_second], 'unpublish').json()['changed'] == 2
    assert counts()[1] == 2
    unpublished = [service.detail('resource', r['source_id']) for r in [first, second]]
    assert batch(unpublished, 'unpublish').json() == {'changed': 0, 'skipped': 2}
    assert batch(unpublished).json()['changed'] == 2
    assert counts()[1] == 4


def test_batch_invalid_later_resource_rolls_back_earlier_publication(client, admin, metadata):
    first = create(admin, metadata)
    second = service.save('resource', 'invalid-level', service.ResourceSave(base_revision=0, metadata={**metadata, 'level': None}), admin['id'])
    before = counts()
    response = client.post(PATH + '/batch', headers=auth_headers(admin), json={'action': 'publish', 'items': [{'source_id': row['source_id'], 'base_revision': row['revision']} for row in [first, second]]})
    assert response.status_code == 409 and '本批未生效' in response.json()['detail']
    assert counts() == before and service.detail('resource', first['source_id'])['status'] == 'draft'


def test_all_transfer_and_batch_endpoints_require_admin(client, admin):
    ordinary = make_user('reader')
    for headers, expected in [({}, 401), (auth_headers(ordinary), 403)]:
        assert client.get(PATH + '/export', headers=headers).status_code == expected
        assert client.post(PATH + '/import', headers=headers, files={'file': ('test.xlsx', b'bad')}).status_code == expected
        assert client.post(PATH + '/batch', headers=headers, json={'action': 'publish', 'items': [{'source_id': 'x', 'base_revision': 1}]}).status_code == expected

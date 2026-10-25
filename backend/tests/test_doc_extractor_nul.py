"""NUL-only extraction cleanup and isolated persistence; parser outputs are synthetic."""
from types import SimpleNamespace
import pytest
from backend.app import doc_extractor,material_files
from backend.app.database import get_db
from .conftest import make_partner


def native_output(monkeypatch,kind,text):
    """Mock only native parser output; preserve the real extract_text assembly."""
    import PyPDF2,docx,pptx
    import docx.table,docx.text.paragraph
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if kind=='pdf':
        pages=[SimpleNamespace(extract_text=lambda:text),SimpleNamespace(extract_text=lambda:'第二页\n甲\t乙')]
        monkeypatch.setattr(PyPDF2,'PdfReader',lambda *_a,**_k:SimpleNamespace(is_encrypted=False,pages=pages))
        return (text+'\n\n第二页\n甲\t乙').strip()
    if kind=='docx':
        class Paragraph:
            def __init__(self,value):self.text=value
        class Table:
            def __init__(self):
                cells=[SimpleNamespace(_tc=object(),iter_inner_content=lambda value=value:iter([Paragraph(value)])) for value in ('表一','表二')]
                self.rows=[SimpleNamespace(cells=cells)]
        monkeypatch.setattr(docx.text.paragraph,'Paragraph',Paragraph)
        monkeypatch.setattr(docx.table,'Table',Table)
        monkeypatch.setattr(docx,'Document',lambda *_:SimpleNamespace(iter_inner_content=lambda:iter([Paragraph(text),Table()])))
        return (text+'\n表一\t表二').strip()
    shapes=[
        SimpleNamespace(shape_type=MSO_SHAPE_TYPE.AUTO_SHAPE,has_text_frame=True,text=text),
        SimpleNamespace(shape_type=MSO_SHAPE_TYPE.TABLE,has_text_frame=False,has_table=True,table=SimpleNamespace(rows=[SimpleNamespace(cells=[SimpleNamespace(text=v,is_spanned=False) for v in ('表一','表二')])]))
    ]
    monkeypatch.setattr(pptx,'Presentation',lambda *_:SimpleNamespace(slides=[SimpleNamespace(shapes=shapes)]))
    return '第 1 页\n'+text+'\n表一\t表二'


def long_text(nul):
    text='  828B2B企业节\r\n表格甲\t表格乙\n中文é🚀\f其他控制字符\x01\n'+'完整正文'*3300+'END_OF_FILE  '
    return '\x00 '+text[:6]+'\x00'+text[6:]+' \x00' if nul else text


@pytest.mark.parametrize('kind',['pdf','docx','pptx'])
@pytest.mark.parametrize('nul',[False,True],ids=['already-clean','contains-nul'])
def test_native_extraction_removes_only_nul_and_preserves_full_text(monkeypatch,kind,nul):
    raw=long_text(nul);before=native_output(monkeypatch,kind,raw)
    result=doc_extractor.extract_text('/tmp/synthetic-native.'+kind,kind)
    assert result==before.replace('\x00','')
    assert '\x00' not in result
    assert len(result)==len(before)-before.count('\x00')
    assert len(result)>13000 and 'END_OF_FILE' in result
    assert '\r\n' in result and '\t' in result and '\f' in result
    assert 'é🚀' in result and '\x01' in result
    if not nul:assert result==before


@pytest.mark.parametrize('kind',['pdf','docx','pptx'])
def test_native_parser_error_still_propagates(monkeypatch,kind):
    import PyPDF2,docx,pptx
    def fail(*_a,**_k):raise RuntimeError('synthetic native parser failure')
    module,name={'pdf':(PyPDF2,'PdfReader'),'docx':(docx,'Document'),'pptx':(pptx,'Presentation')}[kind]
    monkeypatch.setattr(module,name,fail)
    with pytest.raises(RuntimeError,match='synthetic native parser failure'):
        doc_extractor.extract_text('/tmp/synthetic-error.'+kind,kind)


@pytest.mark.parametrize('kind',['txt','markdown','html'])
def test_text_decoding_still_rejects_nul(tmp_path,kind):
    p=tmp_path/('rejected.'+kind);p.write_bytes(b'normal\x00text')
    with pytest.raises(ValueError,match='文件不是可读取的文本'):
        doc_extractor.extract_text(str(p),kind)


@pytest.mark.parametrize('kind',['pdf','docx','pptx'])
@pytest.mark.parametrize('scope',['document','attachment'])
def test_native_nul_result_persists_without_losing_text(client,monkeypatch,kind,scope):
    pid=make_partner()['id'];fid='synthetic-nul-file';cid='synthetic-nul-case'
    expected=native_output(monkeypatch,kind,long_text(True)).replace('\x00','')
    monkeypatch.setattr(material_files,'office_preview',lambda *_:None)
    monkeypatch.setattr(material_files,'process_contribution',lambda *_a,**_k:None)
    table,parent=material_files.TABLES[scope]
    with get_db() as conn:
        assert conn.execute('SELECT current_database()').fetchone()[0] in ('banfei_agent_test','banfei_validation')
        assert conn.execute('SELECT current_schema()').fetchone()[0].startswith('validation_')
        if scope=='attachment':
            conn.execute("INSERT INTO cases (id,partner_id,title,category_id,visible,created_at,updated_at) VALUES (?,?,?,'technical-1',0,'2026-10-25','2026-10-25')",(cid,pid,'合成NUL资料'))
        conn.execute(f'INSERT INTO {table} (id,{parent},filename,file_path,file_type,created_at) VALUES (?,?,?,?,?,?)',
                     (fid,pid if scope=='document' else cid,'合成.'+kind,'/tmp/synthetic-native.'+kind,kind,'2026-10-25'))
    material_files.process(scope,fid)
    with get_db() as conn:
        row=conn.execute(f'SELECT * FROM {table} WHERE id=?',(fid,)).fetchone()
        assert row['extracted_text']==expected
        assert row['processing_status']=='ready' and row['processing_error'] is None
        assert '\x00' not in row['extracted_text'] and len(row['extracted_text'])>13000

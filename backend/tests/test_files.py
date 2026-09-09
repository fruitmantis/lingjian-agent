import os
import io
import uuid
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from pptx import Presentation

from backend.app.database import UPLOADS_DIR, get_db

from .conftest import auth_headers, make_partner, make_user


def docx_bytes(text: str = "validation document") -> bytes:
    output = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(output)
    return output.getvalue()


def pptx_bytes(text: str) -> bytes:
    output = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 3000000, 1000000)
    box.text = text
    presentation.save(output)
    return output.getvalue()


def make_case(partner_id: str) -> str:
    case_id = str(uuid.uuid4())
    with get_db() as conn:
        conn.execute(
            "INSERT INTO cases (id, partner_id, title, description, created_at) VALUES (?, ?, '验证案例', '', ?)",
            (case_id, partner_id, datetime.now(timezone.utc).isoformat()),
        )
    return case_id


def test_file_001_normal_document_upload(client):
    admin = make_user("file_admin", role="admin")
    partner = make_partner()
    response = client.post(
        f"/partners/{partner['id']}/documents", headers=auth_headers(admin),
        files={"file": ("evidence.docx", docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201
    assert "validation document" in response.json()["extracted_text"]
    with get_db() as conn:
        row = conn.execute("SELECT file_path FROM partner_documents WHERE id = ?", (response.json()["id"],)).fetchone()
    assert Path(row["file_path"]).exists()


def test_file_002_upload_size_boundary(client, monkeypatch):
    monkeypatch.setenv("MAX_UPLOAD_SIZE_BYTES", "2048")
    admin = make_user("size_admin", role="admin")
    partner = make_partner()
    case_id = make_case(partner["id"])
    headers = auth_headers(admin)
    below = client.post(
        f"/cases/{case_id}/deliverables", headers=headers,
        files={"file": ("below.bin", b"a" * 2048, "application/octet-stream")},
    )
    assert below.status_code == 201
    above = client.post(
        f"/cases/{case_id}/deliverables", headers=headers,
        files={"file": ("above.bin", b"a" * 2049, "application/octet-stream")},
    )
    assert above.status_code == 413
    assert not any(path.name.endswith("_above.bin") for path in UPLOADS_DIR.glob("*"))


def test_file_003_forged_office_extension_is_rejected(client):
    admin = make_user("signature_admin", role="admin")
    partner = make_partner()
    response = client.post(
        f"/partners/{partner['id']}/documents", headers=auth_headers(admin),
        files={"file": ("malicious.pptx", b"<script>alert(1)</script>", "application/octet-stream")},
    )
    assert response.status_code == 400
    with get_db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM partner_documents").fetchone()[0] == 0
    assert list(UPLOADS_DIR.glob("*")) == []


def test_file_004_missing_disk_file_returns_controlled_responses(client):
    admin = make_user("missing_file_admin", role="admin")
    partner = make_partner()
    document_id = str(uuid.uuid4())
    missing = UPLOADS_DIR / "missing.docx"
    now = datetime.now(timezone.utc).isoformat()
    with get_db() as conn:
        conn.execute(
            """INSERT INTO partner_documents
               (id, partner_id, filename, file_path, file_type, created_at)
               VALUES (?, ?, 'missing.docx', ?, 'docx', ?)""",
            (document_id, partner["id"], str(missing), now),
        )
    headers = auth_headers(admin)
    assert client.get(f"/partners/{partner['id']}/documents/{document_id}/file", headers=headers).status_code == 404
    assert client.get(f"/partners/{partner['id']}/documents/{document_id}/preview", headers=headers).status_code == 404
    assert client.delete(f"/partners/{partner['id']}/documents/{document_id}", headers=headers).status_code == 204


def test_file_005_database_delete_failure_keeps_disk_file(client_no_raise):
    admin = make_user("delete_failure_admin", role="admin")
    partner = make_partner()
    document_id = str(uuid.uuid4())
    file_path = UPLOADS_DIR / "must-remain.docx"
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    file_path.write_bytes(docx_bytes())
    now = datetime.now(timezone.utc).isoformat()
    with get_db() as conn:
        conn.execute(
            """INSERT INTO partner_documents
               (id, partner_id, filename, file_path, file_type, created_at)
               VALUES (?, ?, 'must-remain.docx', ?, 'docx', ?)""",
            (document_id, partner["id"], str(file_path), now),
        )
        if os.environ['DATABASE_URL'].startswith('postgresql'):
            conn.execute("""CREATE FUNCTION validation_document_delete_failure() RETURNS trigger
                LANGUAGE plpgsql AS $$ BEGIN RAISE EXCEPTION 'injected delete failure'; END $$""")
            conn.execute("""CREATE TRIGGER prevent_validation_document_delete BEFORE DELETE ON partner_documents
                FOR EACH ROW EXECUTE FUNCTION validation_document_delete_failure()""")
        else:
            conn.execute(
                """CREATE TRIGGER prevent_validation_document_delete
                   BEFORE DELETE ON partner_documents
                   BEGIN SELECT RAISE(ABORT, 'injected delete failure'); END"""
            )
    response = client_no_raise.delete(
        f"/partners/{partner['id']}/documents/{document_id}", headers=auth_headers(admin),
    )
    assert response.status_code == 500
    assert file_path.exists()
    with get_db() as conn:
        assert conn.execute("SELECT COUNT(*) FROM partner_documents WHERE id = ?", (document_id,)).fetchone()[0] == 1


def test_file_006_pptx_preview_escapes_html(client):
    admin = make_user("preview_admin", role="admin")
    partner = make_partner()
    dangerous = '<script>alert(1)</script><img src=x onerror="alert(2)">'
    upload = client.post(
        f"/partners/{partner['id']}/documents", headers=auth_headers(admin),
        files={"file": ("preview.pptx", pptx_bytes(dangerous), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert upload.status_code == 201
    preview = client.get(
        f"/partners/{partner['id']}/documents/{upload.json()['id']}/preview",
        headers=auth_headers(admin),
    )
    assert preview.status_code == 200
    assert dangerous not in preview.text
    assert "&lt;script&gt;" in preview.text
    assert "default-src 'none'" in preview.headers["content-security-policy"]

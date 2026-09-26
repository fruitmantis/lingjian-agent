from io import BytesIO

from openpyxl import load_workbook

from backend.app import partner_transfer as transfer
from backend.app.database import get_db
from backend.tests.conftest import make_partner, make_user, auth_headers


def save(book):
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def upload(client, headers, book):
    return client.post('/partners/import', headers=headers, files={'file': ('partners.xlsx', save(book))})


def set_cell(book, key, value, row=2):
    book['伙伴信息'].cell(row, list(transfer.COLUMNS).index(key) + 1).value = value


def test_full_partner_roundtrip_preserves_long_profile_pending_and_cases(client):
    headers = auth_headers(make_user('excel-admin', role='admin'))
    make_partner()
    profile = '=画像全文\n' + '完整画像内容。' * 10000 + '\n结尾'
    with get_db() as conn:
        conn.execute("UPDATE partners SET ai_profile=?,intro='简介全文',capabilities='已有能力',industries='金融,待确认行业',service_areas='广东,欧洲,待确认区域',status='disabled',materials_revision=4,profile_materials_revision=3 WHERE id='partner-1'", (profile,))
        conn.execute("INSERT INTO cases (id,partner_id,title,description,created_at) VALUES ('case-kept','partner-1','不能导出或更改的案例','案例原文','2026-09-26')")
    response = client.get('/partners/export', headers=headers)
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book['伙伴信息'].max_row == 2 and book['长文本续文'].max_row > 2
    assert '案例原文' not in str([[cell.value for row in sheet for cell in row] for sheet in book])
    assert book['伙伴信息'].cell(2, list(transfer.COLUMNS).index('ai_profile') + 1).data_type == 's'
    assert upload(client, headers, book).json() == {'created': 0, 'updated': 0, 'unchanged': 1}
    set_cell(book, 'id', 'new-partner')
    set_cell(book, 'name', '新环境伙伴')
    for row in book['长文本续文'].iter_rows(min_row=2):
        row[0].value = 'new-partner'
        row[1].value = '新环境伙伴'
    assert upload(client, headers, book).json()['created'] == 1
    with get_db() as conn:
        imported = dict(conn.execute("SELECT * FROM partners WHERE id='new-partner'").fetchone())
        assert imported['ai_profile'] == profile
        assert imported['industries'] == '金融,待确认行业'
        assert imported['service_areas'] == '广东,欧洲,待确认区域'
        assert imported['status'] == 'disabled'
        assert imported['materials_revision'] != imported['profile_materials_revision']
        assert tuple(conn.execute("SELECT partner_id,description FROM cases WHERE id='case-kept'").fetchone()) == ('partner-1', '案例原文')
        assert conn.execute('SELECT count(*) FROM cases').fetchone()[0] == 1


def test_template_import_name_dedup_and_error_rollback(client):
    headers = auth_headers(make_user('template-admin', role='admin'))
    book = load_workbook(BytesIO(client.get('/partners/template', headers=headers).content))
    assert book['伙伴信息'].max_row == 1 and book['长文本续文'].max_row == 1
    set_cell(book, 'name', '新建伙伴')
    set_cell(book, 'name', '错误伙伴', row=3)
    set_cell(book, 'industries', '任意非标准行业', row=3)
    response = upload(client, headers, book)
    assert response.status_code == 422 and '第3行' in response.json()['detail']
    with get_db() as conn:
        assert conn.execute('SELECT count(*) FROM partners').fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM user_audit_logs WHERE action='partners_imported'").fetchone()[0] == 0
    book['伙伴信息'].delete_rows(3)
    assert upload(client, headers, book).json()['created'] == 1
    assert upload(client, headers, book).json()['unchanged'] == 1
    with get_db() as conn:
        target_id = conn.execute('SELECT id FROM partners').fetchone()[0]
    set_cell(book, 'id', 'different-environment-id')
    set_cell(book, 'ai_profile', '管理员直接导入的画像')
    assert upload(client, headers, book).json()['updated'] == 1
    with get_db() as conn:
        assert tuple(conn.execute('SELECT id,ai_profile FROM partners').fetchone()) == (target_id, '管理员直接导入的画像')
        assert conn.execute('SELECT count(*) FROM partners').fetchone()[0] == 1


def test_templates_and_transfer_require_admin_and_resource_template_is_importable(client):
    admin = auth_headers(make_user('access-admin', role='admin'))
    reader = auth_headers(make_user('access-reader'))
    for headers, status in [({}, 401), (reader, 403)]:
        for url in ['/partners/template', '/partners/export', '/admin/enablement/resources/template']:
            assert client.get(url, headers=headers).status_code == status
        assert client.post('/partners/import', headers=headers, files={'file': ('test.xlsx', b'bad')}).status_code == status
    response = client.get('/admin/enablement/resources/template', headers=admin)
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book['课程与实验'].max_row == 1 and book['岗位与专区'].max_row > 1
    from backend.app.resource_transfer import COLUMNS, append_text
    values = {'resource_type': '课程', 'title': '模板课程', 'summary': '课程简介', 'level': '基础',
              'source_url': 'https://example.com/course', 'system_visible': '是', 'model_allowed': '否', 'partner_allowed': '否'}
    append_text(book['课程与实验'], [values.get(key) for key in COLUMNS])
    response = client.post('/admin/enablement/resources/import', headers=admin, files={'file': ('resources.xlsx', save(book))})
    assert response.status_code == 200 and response.json()['created'] == 1


def test_invalid_continuation_and_formula_do_not_overwrite_partner(client):
    headers = auth_headers(make_user('validation-admin', role='admin'))
    make_partner()
    with get_db() as conn:
        conn.execute("UPDATE partners SET ai_profile=? WHERE id='partner-1'", ('全文' * 40000,))
    original = client.get('/partners/export', headers=headers).content
    book = load_workbook(BytesIO(original))
    book['长文本续文'].delete_rows(2)
    assert upload(client, headers, book).status_code == 422
    book = load_workbook(BytesIO(original))
    set_cell(book, 'intro', '=1+1')
    assert upload(client, headers, book).status_code == 422
    with get_db() as conn:
        assert conn.execute("SELECT ai_profile FROM partners WHERE id='partner-1'").fetchone()[0] == '全文' * 40000

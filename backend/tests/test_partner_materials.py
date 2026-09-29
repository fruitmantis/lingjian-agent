"""Current contract: six native document types, cached extraction and live case visibility."""
import io
import json
from pathlib import Path
import pytest
from PyPDF2 import PdfWriter
from docx import Document
from pptx import Presentation
from backend.app.database import get_db
from backend.app import material_files,doc_extractor
from .conftest import auth_headers,make_partner,make_user
from .test_files import docx_bytes,pptx_bytes

@pytest.fixture
def access(client):
    admin=auth_headers(make_user('material-admin',role='admin'))
    ordinary=auth_headers(make_user('material-reader'))
    partner=make_partner()
    return admin,ordinary,partner['id']

def upload(client,access,name,data,initialize=False):
    admin,_,pid=access
    r=client.post(f'/partners/{pid}/documents',headers=admin,files={'file':(name,data)},data={'initialize_profile':str(initialize).lower()})
    assert r.status_code==201,r.text
    fid=r.json()['id']
    with get_db() as conn: return dict(conn.execute('SELECT * FROM partner_documents WHERE id=?',(fid,)).fetchone())

@pytest.mark.parametrize('name,data,expected',[
    ('note.txt','中文常用编码'.encode('gb18030'),'中文常用编码'),
    ('note.md',b'# Heading\n\n- native text\n\n![image alt](https://outside.test/a.png)','Heading'),
    ('image.markdown',b'# Heading\n\n<img src="https://outside.test/a" alt="LEAK_IMAGE">\n\n<p>native</p>','native'),
    ('note.markdown',b'| A | B |\n|---|---|\n| one | two |','one'),
    ('note.html',b'<html><body><h1>Title</h1><p>native text</p><script>LEAK_SCRIPT</script><img alt="LEAK_IMAGE" src="https://outside.test/a"></body></html>','native text'),
    ('note.htm',b'<p>Native HTML</p>','Native HTML'),
    ('note.docx',docx_bytes('native docx'),'native docx'),
    ('note.pptx',pptx_bytes('native pptx'),'native pptx'),
])
def test_six_formats_native_text_and_private_original(client,access,name,data,expected):
    row=upload(client,access,name,data)
    assert row['processing_status']=='ready',row['processing_error']
    assert expected in row['extracted_text']
    assert not row['preview_error'],row['preview_error']
    assert 'LEAK_SCRIPT' not in row['extracted_text'] and 'LEAK_IMAGE' not in row['extracted_text']
    assert Path(row['file_path']).read_bytes()==data
    admin,user,pid=access;path=f'/partners/{pid}/documents/{row["id"]}'
    assert client.get(path+'/file',headers=admin).content==data
    assert client.get(path+'/preview',headers=admin).status_code==200
    assert client.get(path+'/file',headers=user).status_code==403
    response=client.get(f'/partners/{pid}/documents',headers=admin).json()[0]
    assert not {'file_path','preview_path'}&response.keys()

@pytest.mark.parametrize('name',['old.doc','old.ppt','image.png','sheet.xlsx','archive.zip','app.exe'])
def test_unsupported_formats_rejected(client,access,name):
    admin,_,pid=access
    r=client.post(f'/partners/{pid}/documents',headers=admin,files={'file':(name,b'fake')})
    assert r.status_code==400
    with get_db() as conn: assert conn.execute('SELECT count(*) FROM partner_documents').fetchone()[0]==0

@pytest.mark.parametrize('name,data',[('rename.pdf',b'not PDF'),('rename.docx',pptx_bytes('wrong package')),('rename.txt',docx_bytes()),('rename.html',b'not HTML')])
def test_extension_rename_does_not_bypass_validation(client,access,name,data):
    admin,_,pid=access
    assert client.post(f'/partners/{pid}/documents',headers=admin,files={'file':(name,data)}).status_code==400

def test_scan_pdf_saved_previewable_but_not_profile_input(client,access):
    pdf=PdfWriter();pdf.add_blank_page(100,100);b=io.BytesIO();pdf.write(b)
    row=upload(client,access,'scan.pdf',b.getvalue())
    assert row['processing_status']=='empty' and not row['extracted_text']
    assert client.get(f'/partners/{access[2]}/documents/{row["id"]}/preview',headers=access[0]).content.startswith(b'%PDF-')

def test_text_saved_in_full_and_cached_without_reextract(client,access,monkeypatch):
    text='完整资料'*10000+'END_OF_FILE'
    row=upload(client,access,'long.txt',text.encode())
    assert row['extracted_text']==text
    monkeypatch.setattr(material_files,'extract_text',lambda *_:pytest.fail('Viewing must not extract again'))
    assert client.get(f'/partners/{access[2]}/documents/{row["id"]}/preview',headers=access[0]).text==text

def test_docx_tables_pptx_tables_images_ignored(tmp_path):
    doc=Document();doc.add_paragraph('正文');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='表一';table.cell(0,1).text='表二'
    path=tmp_path/'table.docx';doc.save(path);text=doc_extractor.extract_text(str(path),'docx')
    assert all(v in text for v in ['正文','表一','表二'])
    ppt=Presentation();s=ppt.slides.add_slide(ppt.slide_layouts[6]);t=s.shapes.add_table(1,2,0,0,2000000,1000000).table;t.cell(0,0).text='甲';t.cell(0,1).text='乙'
    path=tmp_path/'table.pptx';ppt.save(path);assert '甲\t乙' in doc_extractor.extract_text(str(path),'pptx')

def test_safe_preview_has_no_active_or_remote_content(client,access):
    row=upload(client,access,'page.html',b'<p onclick="alert(1)">Text</p><script>SECRET_SCRIPT</script><iframe src="https://evil.test"></iframe><img src="https://evil.test/a"><a href="javascript:alert(1)">Link</a>')
    r=client.get(f'/partners/{access[2]}/documents/{row["id"]}/preview',headers=access[0])
    assert r.status_code==200
    assert all(v not in r.text for v in ['onclick','javascript:','https://evil','SECRET_SCRIPT','<iframe','<img'])
    assert "default-src 'none'" in r.headers['content-security-policy']

def test_direct_docx_profile_adoption_and_failure_preserve_old(client,access,monkeypatch):
    from backend.app.routers import profile
    monkeypatch.setattr(profile,'chat_completion',lambda *_a,**_k:pytest.fail('Import must not call AI'))
    text='原样采用的完整画像'*12000
    assert len(text)>100000  # Direct DOCX adoption does not use the AI input budget.
    row=upload(client,access,'profile.docx',docx_bytes(text),True)
    with get_db() as conn: assert conn.execute('SELECT ai_profile FROM partners WHERE id=?',(access[2],)).fetchone()[0]==text
    monkeypatch.setattr(material_files,'extract_text',lambda *_:(_ for _ in ()).throw(ValueError('synthetic extraction failure')))
    failed=upload(client,access,'fail.docx',docx_bytes('new content'),True)
    assert failed['processing_status']=='failed' and Path(failed['file_path']).exists()
    assert 'synthetic extraction failure' in failed['processing_error']
    with get_db() as conn: assert conn.execute('SELECT ai_profile FROM partners WHERE id=?',(access[2],)).fetchone()[0]==text
    assert client.get(f'/partners/{access[2]}',headers=access[1]).json()['ai_profile'] is None
    assert '原样采用' not in client.get('/enablement/context',params={'partner_id':access[2]},headers=access[1]).text

def test_preview_failure_retains_text_and_retry_preserves_profile(client,access,monkeypatch):
    monkeypatch.setattr(material_files,'office_preview',lambda *_:(_ for _ in ()).throw(RuntimeError('synthetic converter failure')))
    row=upload(client,access,'preview.docx',docx_bytes('retained text'))
    assert row['processing_status']=='ready' and row['extracted_text']=='retained text'
    assert 'synthetic converter failure' in row['preview_error']

def test_manual_profile_full_input_budget_and_failed_generation(client,access,monkeypatch):
    from backend.app.routers import profile
    text='完整资料'*2000+'END_OF_FILE';upload(client,access,'full.txt',text.encode())
    messages=[]
    def complete(msg,**kwargs):
        messages.append(msg)
        return '{}' if len(messages)%2 else '新的伙伴画像'
    monkeypatch.setattr(profile,'chat_completion',complete)
    response=client.post(f'/partners/{access[2]}/profile',headers=access[0])
    assert response.status_code==200,response.text
    assert all(text in m[1]['content'] for m in messages)
    with get_db() as conn:
        p=dict(conn.execute('SELECT * FROM partners WHERE id=?',(access[2],)).fetchone());assert p['materials_revision']==p['profile_materials_revision']
    monkeypatch.setenv('BANFEI_PROFILE_INPUT_MAX_CHARS','1000');messages.clear()
    response=client.post(f'/partners/{access[2]}/profile',headers=access[0]);assert response.status_code==422 and not messages
    with get_db() as conn: assert conn.execute('SELECT ai_profile FROM partners WHERE id=?',(access[2],)).fetchone()[0]=='新的伙伴画像'
    monkeypatch.delenv('BANFEI_PROFILE_INPUT_MAX_CHARS',raising=False)
    monkeypatch.setattr(profile,'chat_completion',lambda *_a,**_k:(_ for _ in ()).throw(RuntimeError('synthetic model failure')))
    assert client.post(f'/partners/{access[2]}/profile',headers=access[0]).status_code==502
    with get_db() as conn: assert conn.execute('SELECT ai_profile FROM partners WHERE id=?',(access[2],)).fetchone()[0]=='新的伙伴画像'


@pytest.mark.parametrize('context_chars',[100000,100001])
def test_default_profile_budget_boundary_keeps_full_input_and_old_profile(client,access,monkeypatch,context_chars):
    from backend.app.routers import profile
    monkeypatch.delenv('BANFEI_PROFILE_INPUT_MAX_CHARS',raising=False)
    source='首尾';row=upload(client,access,'boundary.txt',source.encode())
    messages=[]
    def complete(msg,**kwargs):
        messages.append(msg)
        return '{}' if len(messages)%2 else '合成边界画像'
    monkeypatch.setattr(profile,'chat_completion',complete)
    endpoint=f'/partners/{access[2]}/profile'
    assert client.post(endpoint,headers=access[0]).status_code==200
    overhead=len(messages[0][1]['content'])-len(source)
    source='首'+'字'*(context_chars-overhead-2)+'尾'
    with get_db() as conn:
        conn.execute('UPDATE partner_documents SET extracted_text=? WHERE id=?',(source,row['id']))
        before=dict(conn.execute('SELECT * FROM partners WHERE id=?',(access[2],)).fetchone())
    messages.clear()
    response=client.post(endpoint,headers=access[0])
    if context_chars==100000:
        assert response.status_code==200,response.text
        assert len(messages)==2
        assert all(len(m[1]['content'])==context_chars and source in m[1]['content'] for m in messages)
    else:
        assert response.status_code==422
        assert '共 100001 字' in response.json()['detail'] and '上限 100000 字' in response.json()['detail']
        assert not messages
        with get_db() as conn:
            assert dict(conn.execute('SELECT * FROM partners WHERE id=?',(access[2],)).fetchone())==before

def test_case_visibility_files_categories_and_live_reference(client,access):
    admin,user,pid=access
    categories=client.get('/cases/categories',headers=user).json()
    assert len(categories)==3 and all(len(c['children'])==6 for c in categories)
    payload={'partner_id':pid,'title':'案例一','category_id':'technical-3','description':'静态正文'}
    r=client.post('/cases',headers=admin,json=payload);assert r.status_code==201,r.text
    cid=r.json()['id'];path=f'/cases/{cid}'
    upload_response=client.post(path+'/deliverables',headers=admin,files={'file':('case.txt','案例附件正文'.encode())});assert upload_response.status_code==201,upload_response.text
    fid=upload_response.json()['id'];file_url=path+f'/deliverables/{fid}/file'
    assert client.get(path,headers=user).status_code==404 and client.get(file_url,headers=user).status_code==404
    assert client.get('/enablement/resources?source_type=case',headers=user).json()['total']==0
    assert client.patch(path+'/visibility',headers=admin,json={'visible':True}).status_code==200
    assert client.get(file_url,headers=user).text=='案例附件正文'
    assert client.get('/enablement/resources?source_type=case&category_group=technical',headers=user).json()['total']==1
    assert client.get('/enablement/resources?source_type=case&category_id=delivery-1',headers=user).json()['total']==0
    payload.update(title='立即更新',visible=True)
    assert client.put(path,headers=admin,json=payload).status_code==200
    detail=client.get(f'/enablement/resources/case/{cid}?source_version=99',headers=user).json()
    assert detail['title']=='立即更新' and len(detail['files'])==1
    assert not {'reviews','versions','source_url','draft_json','authorization_epoch'}&detail.keys()
    assert client.patch(path+'/visibility',headers=admin,json={'visible':False}).status_code==200
    assert client.get(file_url,headers=user).status_code==404
    assert client.get(f'/enablement/resources/case/{cid}',headers=user).status_code==404
    assert client.get(file_url,headers=admin).status_code==200
    with get_db() as conn:
        assert conn.execute('SELECT count(*) FROM enablement_audit_events WHERE source_kind=?',('case',)).fetchone()[0]==0
    assert client.delete(path,headers=admin).status_code==204

def test_case_requires_partner_valid_category_and_admin(client,access):
    admin,user,pid=access
    payload={'partner_id':pid,'title':'case','category_id':'technical-1'}
    assert client.post('/cases',headers=user,json=payload).status_code==403
    assert client.post('/cases',headers=admin,json={**payload,'partner_id':'missing'}).status_code==404
    assert client.post('/cases',headers=admin,json={**payload,'category_id':'invented'}).status_code==422
    assert client.put('/admin/cases/any/sharing',headers=admin,json={}).status_code==404


def test_native_text_survives_images_without_ocr(tmp_path):
    from PIL import Image
    image=tmp_path/'picture.png';Image.new('RGB',(100,100),'white').save(image)
    doc=Document();doc.add_paragraph('native only');doc.add_picture(str(image))
    doc_path=tmp_path/'mixed.docx';doc.save(doc_path)
    assert doc_extractor.extract_text(str(doc_path),'docx')=='native only'
    ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[6]);slide.shapes.add_picture(str(image),0,0)
    ppt_path=tmp_path/'image.pptx';ppt.save(ppt_path)
    assert doc_extractor.extract_text(str(ppt_path),'pptx')==''


def test_pdf_existing_text_layer_is_extracted_without_ocr(client,access):
    from PyPDF2.generic import DictionaryObject,NameObject,DecodedStreamObject
    writer=PdfWriter();writer.add_blank_page(200,100);page=writer.pages[0]
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 10 50 Td (Native PDF text) Tj ET');page[NameObject('/Contents')]=stream
    buf=io.BytesIO();writer.write(buf)
    row=upload(client,access,'native.pdf',buf.getvalue())
    assert row['processing_status']=='ready' and 'Native PDF text' in row['extracted_text']


@pytest.mark.parametrize('newer_profile',[False,True])
def test_failed_initial_profile_can_retry_without_overwriting_newer_profile(client,access,monkeypatch,newer_profile):
    original=material_files.extract_text
    monkeypatch.setattr(material_files,'extract_text',lambda *_:(_ for _ in ()).throw(ValueError('temporary parser failure')))
    row=upload(client,access,'retry.docx',docx_bytes('导入画像正文'),True)
    assert row['processing_status']=='failed'
    if newer_profile:
        with get_db() as conn:
            conn.execute('UPDATE partners SET ai_profile=?,profile_updated_at=? WHERE id=?',('更新后的画像','2999-01-01',access[2]))
    monkeypatch.setattr(material_files,'extract_text',original)
    r=client.post(f'/partners/{access[2]}/documents/{row["id"]}/retry',headers=access[0]);assert r.status_code==200
    with get_db() as conn:
        assert conn.execute('SELECT processing_status FROM partner_documents WHERE id=?',(row['id'],)).fetchone()[0]=='ready'
        assert conn.execute('SELECT ai_profile FROM partners WHERE id=?',(access[2],)).fetchone()[0]==('更新后的画像' if newer_profile else '导入画像正文')
